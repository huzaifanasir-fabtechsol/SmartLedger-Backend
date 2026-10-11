"""
Bulk expense + transaction import from an .xlsx file.

The work is split in three steps so nothing is written until the whole file is known to be good:

  1. read_sheet()       xlsx -> plain rows (verifies the columns)
  2. build_payload()    validates every row and returns a JSON-safe payload, or raises
                        ImportFileError listing every problem. Nothing touches the database.
  3. process_import_job()  (Celery task) writes the payload inside ONE database transaction,
                        so a job is either imported completely or not at all.

Re-importing the same file is safe: rows are matched one-to-one against what already exists.
"""
import logging
import re
from collections import defaultdict, deque
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.conf import settings
from django.core.cache import cache
from django.db import IntegrityError, connection, transaction as db_transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.expense.models import Expense, ExpenseCategory, ExpenseImportJob
from apps.revenue.models import Transaction

logger = logging.getLogger(__name__)

CHUNK_SIZE = 500
MAX_REPORTED_ERRORS = 200
MAX_REPORTED_WARNINGS = 50
MAX_SCANNED_ROWS = 200000

TWO_PLACES = Decimal('0.01')
MAX_TRANSACTION_AMOUNT = Decimal('9999999999.99')   # Transaction.withdraw / deposit: 12 digits
MAX_EXPENSE_AMOUNT = Decimal('99999999.99')         # Expense.amount: 10 digits

# (field, label shown to the user, accepted header spellings after normalisation)
COLUMNS = [
    ('date', 'Date', {'date'}),
    ('transaction_id', 'Transaction ID', {'transaction id', 'transection id', 'transactionid', 'transectionid', 'txn id'}),
    ('withdraw', 'Withdraw Amount', {'withdraw amount', 'withdraw', 'withdrawal amount', 'withdrawal'}),
    ('deposit', 'Deposit Amount', {'deposit amount', 'deposit', 'deposite amount', 'deposite', 'depostie amount'}),
    ('description', 'Description', {'description'}),
    ('notes', 'Notes', {'notes', 'note'}),
    ('title', 'Title', {'title'}),
    ('category', 'Category', {'category', 'catagory'}),
    ('tax_category', 'Tax Way', {'tax way', 'tax type', 'tax category', 'tax catagory'}),
    ('tax_rate', 'Tax Rate', {'tax rate', 'tax percent', 'tax'}),
]
COLUMN_LABELS = {field: label for field, label, _ in COLUMNS}
HEADER_ALIASES = {alias: field for field, _, aliases in COLUMNS for alias in aliases}

TAX_WAYS = {
    'inclusive': 'inclusive', 'tax inclusive': 'inclusive', 'incl': 'inclusive',
    'exclusive': 'exclusive', 'tax exclusive': 'exclusive', 'excl': 'exclusive',
}

DATE_FORMATS = (
    '%Y-%m-%d', '%Y/%m/%d', '%d-%m-%Y', '%d/%m/%Y',
    '%m/%d/%Y', '%Y.%m.%d', '%d.%m.%Y', '%Y%m%d',
)


class ImportFileError(Exception):
    """The uploaded file cannot be imported. `row_errors` lists every offending row."""

    def __init__(self, message, row_errors=None):
        super().__init__(message)
        self.message = message
        self.row_errors = row_errors or []


class ImportAborted(Exception):
    """Raised inside the import transaction to roll everything back."""


# --------------------------------------------------------------------------------------
# Step 1: read the sheet
# --------------------------------------------------------------------------------------

def _normalize_header(value):
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).split())


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_sheet(file_obj):
    """Return [(sheet_row_number, {field: raw value})] for every non-blank data row."""
    unreadable = ImportFileError('Could not read the file. Please upload a valid .xlsx file.')
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        raise unreadable

    try:
        return _read_rows(workbook)
    except ImportFileError:
        raise
    except Exception:
        logger.exception('Could not read expense import workbook')
        raise unreadable
    finally:
        workbook.close()


def _read_rows(workbook):
    sheet = workbook.active
    if sheet is None:
        raise ImportFileError('The uploaded file has no worksheet.')
    # Some generators write a wrong sheet size; without this rows could be silently skipped.
    sheet.reset_dimensions()

    col_map = None
    rows = []
    for row_number, cells in enumerate(sheet.iter_rows(), start=1):
        if row_number > MAX_SCANNED_ROWS:
            raise ImportFileError('The sheet is too large. Remove unused rows and try again.')
        if all(_text(cell.value) == '' for cell in cells):
            continue

        if col_map is None:
            col_map = _map_columns(cells)
            continue

        record = {}
        for field, idx in col_map.items():
            cell = cells[idx] if idx < len(cells) else None
            value = cell.value if cell is not None else None
            if field == 'tax_rate' and isinstance(value, (int, float)) and not isinstance(value, bool):
                # A cell formatted as a percentage stores 10% as 0.1.
                if '%' in (getattr(cell, 'number_format', None) or ''):
                    value = Decimal(str(value)) * 100
            record[field] = value
        rows.append((row_number, record))

        if len(rows) > settings.EXPENSE_IMPORT_MAX_ROWS:
            raise ImportFileError(
                f'The file has more than {settings.EXPENSE_IMPORT_MAX_ROWS} rows. Split it into smaller files.'
            )

    if col_map is None:
        raise ImportFileError('The uploaded file is empty.')
    if not rows:
        raise ImportFileError('The file has a header row but no data rows.')
    return rows


def _map_columns(header_cells):
    col_map = {}
    duplicates = []
    for idx, cell in enumerate(header_cells):
        field = HEADER_ALIASES.get(_normalize_header(cell.value))
        if not field:
            continue
        if field in col_map:
            duplicates.append(COLUMN_LABELS[field])
        else:
            col_map[field] = idx

    missing = [label for field, label, _ in COLUMNS if field not in col_map]
    problems = []
    if missing:
        problems.append(f"Missing column(s): {', '.join(missing)}")
    if duplicates:
        problems.append(f"Duplicate column(s): {', '.join(sorted(set(duplicates)))}")
    if problems:
        expected = ', '.join(label for _, label, _ in COLUMNS)
        raise ImportFileError(f"{'. '.join(problems)}. Expected columns: {expected}.")
    return col_map


# --------------------------------------------------------------------------------------
# Step 2: validate everything and build the payload
# --------------------------------------------------------------------------------------

def _parse_date(value):
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = _text(value).split(' ')[0]
        parsed = None
        for fmt in DATE_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
    if parsed is None or not 1990 <= parsed.year <= 2100:
        return None
    return parsed


def _parse_decimal(value):
    """Return a non-negative Decimal rounded to 2 places, or None when the value is not usable."""
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        number = value
    else:
        text = _text(value) if not isinstance(value, float) else repr(value)
        text = re.sub(r'[,\s¥￥$]', '', text)
        try:
            number = Decimal(text)
        except (InvalidOperation, ValueError):
            return None
    if not number.is_finite() or number < 0:
        return None
    rounded = number.quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
    # Tolerate float noise (100.10000000001) but not a real third decimal place.
    if abs(number - rounded) > Decimal('0.000001'):
        return None
    return rounded


def _parse_amount(value):
    """Blank and '-' mean zero. Returns None for anything that is not a valid amount."""
    if _text(value) in ('', '-'):
        return Decimal('0.00')
    return _parse_decimal(value)


def _category_key(name):
    return ' '.join(name.split()).casefold()


def split_tax(amount, tax_percent, tax_category):
    """
    Return (base amount, tax amount). Mirrors ExpenseSerializer.validate so imported
    expenses carry exactly the numbers a manually entered one would.
    """
    if tax_percent <= 0:
        return amount, None
    rate = tax_percent / Decimal('100')
    if tax_category == 'inclusive':
        base = (amount / (Decimal('1') + rate)).quantize(TWO_PLACES)
        return base, (amount - base).quantize(TWO_PLACES)
    return amount, (amount * rate).quantize(TWO_PLACES)


def _validate_row(record):
    """Return (clean row, category name, [errors]) for one sheet row."""
    errors = []

    parsed_date = _parse_date(record.get('date'))
    if _text(record.get('date')) == '':
        errors.append('Date is empty')
    elif parsed_date is None:
        errors.append(f"Invalid Date '{_text(record.get('date'))}'")

    transaction_id = _text(record.get('transaction_id'))
    if not transaction_id:
        errors.append('Transaction ID is empty')
    elif len(transaction_id) > 500:
        errors.append('Transaction ID is longer than 500 characters')

    withdraw = _parse_amount(record.get('withdraw'))
    if withdraw is None:
        errors.append(f"Invalid Withdraw Amount '{_text(record.get('withdraw'))}'")
    deposit = _parse_amount(record.get('deposit'))
    if deposit is None:
        errors.append(f"Invalid Deposit Amount '{_text(record.get('deposit'))}'")
    if withdraw is not None and deposit is not None:
        if withdraw == 0 and deposit == 0:
            errors.append('Either Withdraw Amount or Deposit Amount is required')
        elif withdraw > 0 and deposit > 0:
            errors.append('Only one of Withdraw Amount or Deposit Amount can be filled')
        elif max(withdraw, deposit) > MAX_TRANSACTION_AMOUNT:
            errors.append('Amount is too large')

    description = _text(record.get('description'))
    if not description:
        errors.append('Description is empty')
    elif len(description) > 500:
        errors.append('Description is longer than 500 characters')

    notes = _text(record.get('notes'))
    if not notes:
        errors.append('Notes is empty')

    clean = {
        'date': parsed_date.isoformat() if parsed_date else None,
        'transaction_id': transaction_id,
        'withdraw': str(withdraw),
        'deposit': str(deposit),
        'description': description,
        'notes': notes,
        'expense': None,
    }

    # A deposit never becomes an expense, so the expense columns only matter on withdraw rows.
    if not (withdraw and withdraw > 0 and not deposit):
        return clean, None, errors

    title = _text(record.get('title'))
    if not title:
        errors.append('Title is empty')
    elif len(title) > 200:
        errors.append('Title is longer than 200 characters')

    category = ' '.join(_text(record.get('category')).split())
    if not category:
        errors.append('Category is empty')
    elif len(category) > 100:
        errors.append('Category is longer than 100 characters')

    tax_way_text = _text(record.get('tax_category'))
    tax_category = TAX_WAYS.get(' '.join(re.sub(r'[-_]+', ' ', tax_way_text).lower().split()))
    if not tax_way_text:
        errors.append('Tax Way is empty')
    elif tax_category is None:
        errors.append(f"Invalid Tax Way '{tax_way_text}' (use Inclusive or Exclusive)")

    tax_rate_raw = record.get('tax_rate')
    tax_rate_text = _text(tax_rate_raw)
    tax_percent = None
    if not tax_rate_text:
        errors.append('Tax Rate is empty')
    else:
        tax_percent = _parse_decimal(tax_rate_text.rstrip('%') if isinstance(tax_rate_raw, str) else tax_rate_raw)
        if tax_percent is None or tax_percent > 100:
            errors.append(f"Invalid Tax Rate '{tax_rate_text}' (use a number from 0 to 100)")
            tax_percent = None
        elif 0 < tax_percent < 1:
            errors.append(f"Tax Rate '{tax_rate_text}' looks like a fraction; enter a percentage such as 10 for 10%")
            tax_percent = None

    if withdraw > MAX_EXPENSE_AMOUNT:
        errors.append(f'Withdraw Amount is larger than the maximum expense amount ({MAX_EXPENSE_AMOUNT})')

    if not errors:
        amount, tax_amount = split_tax(withdraw, tax_percent, tax_category)
        clean['expense'] = {
            'title': title,
            'category': _category_key(category),
            'tax_category': tax_category,
            'tax_percent': str(tax_percent),
            'amount': str(amount),
            'tax_amount': str(tax_amount) if tax_amount is not None else None,
        }
    return clean, category, errors


def build_payload(file_obj, user):
    """
    Read and validate the whole file. Returns (payload, summary).
    Raises ImportFileError if a single cell anywhere is wrong.
    """
    sheet_rows = read_sheet(file_obj)

    row_errors = []
    rows = []
    file_categories = {}      # key -> spelling used in the file (first occurrence)
    category_rows = defaultdict(list)
    for row_number, record in sheet_rows:
        clean, category, errors = _validate_row(record)
        if errors:
            row_errors.append(f"Row {row_number}: {'; '.join(errors)}")
            continue
        clean['row'] = row_number
        rows.append(clean)
        if category:
            key = _category_key(category)
            file_categories.setdefault(key, category)
            category_rows[key].append(row_number)

    # Categories: reuse the user's own (case / spacing insensitive), create the rest.
    own = {}
    for name in ExpenseCategory.objects.filter(user=user).order_by('-id').values_list('name', flat=True):
        own[_category_key(name)] = name
    new_names = {key: name for key, name in file_categories.items() if key not in own}
    if new_names:
        # ExpenseCategory.name is unique across all users, so a name taken elsewhere cannot be created.
        taken = set(
            ExpenseCategory.objects.filter(name__in=list(new_names.values())).values_list('name', flat=True)
        )
        for key, name in new_names.items():
            if name in taken:
                listed = ', '.join(str(n) for n in category_rows[key][:10])
                row_errors.append(
                    f"Category '{name}' (row {listed}) cannot be created because the name is already "
                    f"used by another account. Use a different category name."
                )

    if row_errors:
        raise ImportFileError(
            f'Validation failed for {len(row_errors)} problem(s). Nothing was imported. '
            f'Fix the file and upload it again.',
            row_errors,
        )

    payload = {
        'categories': {key: own.get(key, name) for key, name in file_categories.items()},
        'rows': rows,
    }
    expense_rows = sum(1 for row in rows if row['expense'])
    summary = {
        'total_rows': len(rows),
        'expense_rows': expense_rows,
        'deposit_rows': len(rows) - expense_rows,
        'new_categories': sorted(new_names.values()),
    }
    return payload, summary


# --------------------------------------------------------------------------------------
# Step 3: write the payload (runs in the Celery worker)
# --------------------------------------------------------------------------------------

def _progress_key(job_id):
    return f'expense_import_progress:{job_id}'


class _ProgressReporter:
    """
    Publishes live progress to the cache. The import runs in one database transaction, so
    writes to the job row would stay invisible to the API until it commits.
    """

    def __init__(self, job_id):
        self.key = _progress_key(job_id)
        self.enabled = True

    def __call__(self, processed):
        if not self.enabled:
            return
        try:
            cache.set(self.key, processed, timeout=settings.EXPENSE_IMPORT_STALE_AFTER)
        except Exception:
            # Progress is cosmetic; never let a cache outage break (or slow down) an import.
            self.enabled = False


def _insert_returning_ids(model, objs):
    if not objs:
        return
    if connection.features.can_return_rows_from_bulk_insert:
        model.objects.bulk_create(objs, batch_size=CHUNK_SIZE)
    else:
        for obj in objs:
            obj.save(force_insert=True)


def _resolve_categories(user_id, wanted):
    """Return ({key: category id}, [created names]) creating whatever does not exist yet."""
    def load():
        found = {}
        for pk, name in ExpenseCategory.objects.filter(user_id=user_id).order_by('-id').values_list('id', 'name'):
            found[_category_key(name)] = pk
        return found

    ids = load()
    missing = [name for key, name in wanted.items() if key not in ids]
    if missing:
        try:
            ExpenseCategory.objects.bulk_create(
                [ExpenseCategory(user_id=user_id, name=name) for name in missing]
            )
        except IntegrityError:
            raise ImportAborted('A new category could not be created because its name is already in use.')
        ids = load()
    return {key: ids[key] for key in wanted}, sorted(missing)


def _recalculate_balances(user_id, account_id, from_date):
    """Same running-balance rule as TransactionViewSet._recalculate_balances."""
    account_txs = Transaction.objects.filter(user_id=user_id, company_account_id=account_id)
    previous = (
        account_txs.filter(date__lt=from_date).order_by('-date', '-id').values_list('balance', flat=True).first()
    )
    running = previous if previous is not None else Decimal('0.00')

    updates = []
    following = account_txs.filter(date__gte=from_date).order_by('date', 'id').only('id', 'deposit', 'withdraw', 'balance')
    for tx in following.iterator(chunk_size=2000):
        running = running + tx.deposit - tx.withdraw
        if tx.balance != running:
            tx.balance = running
            updates.append(tx)
    if updates:
        Transaction.objects.bulk_update(updates, ['balance'], batch_size=CHUNK_SIZE)


def run_import(job, on_progress=None):
    """
    Write job.payload. Must be called inside a transaction; raises on any inconsistency so the
    caller's transaction rolls back and nothing is left half-imported.
    """
    user_id = job.user_id
    account_id = job.company_account_id
    rows = job.payload['rows']
    for row in rows:
        row['date'] = date.fromisoformat(row['date'])
        row['withdraw'] = Decimal(row['withdraw'])
        row['deposit'] = Decimal(row['deposit'])

    category_ids, created_categories = _resolve_categories(user_id, job.payload['categories'])

    # Existing transactions that could match, grouped by (date, transaction id, withdraw, deposit).
    # Each one can be claimed by a single file row, so genuinely repeated bank lines stay separate
    # and re-importing the same file creates nothing.
    file_dates = [row['date'] for row in rows]
    wanted_ids = {(row['date'], row['transaction_id']) for row in rows}
    pool = defaultdict(deque)
    existing = Transaction.objects.filter(
        user_id=user_id,
        company_account_id=account_id,
        date__range=(min(file_dates), max(file_dates)),
    ).order_by('id').values_list('id', 'date', 'transaction_id', 'withdraw', 'deposit')
    for pk, tx_date, tx_id, withdraw, deposit in existing.iterator(chunk_size=2000):
        tx_id = (tx_id or '').strip()
        if (tx_date, tx_id) in wanted_ids:
            pool[(tx_date, tx_id, withdraw, deposit)].append(pk)

    counts = {
        'transactions_created': 0,
        'transactions_existing': 0,
        'expenses_created': 0,
        'expenses_existing': 0,
        'deposit_rows': 0,
    }
    warnings = []
    earliest_new_date = None

    for start in range(0, len(rows), CHUNK_SIZE):
        chunk = rows[start:start + CHUNK_SIZE]

        new_transactions = []
        matched = []   # (row, existing transaction id or new Transaction)
        for row in chunk:
            candidates = pool.get((row['date'], row['transaction_id'], row['withdraw'], row['deposit']))
            if candidates:
                matched.append((row, candidates.popleft()))
                counts['transactions_existing'] += 1
                continue
            tx = Transaction(
                user_id=user_id,
                company_account_id=account_id,
                date=row['date'],
                transaction_id=row['transaction_id'],
                withdraw=row['withdraw'],
                deposit=row['deposit'],
                balance=Decimal('0.00'),
                description=row['description'],
                notes=row['notes'],
            )
            new_transactions.append(tx)
            matched.append((row, tx))
            if earliest_new_date is None or row['date'] < earliest_new_date:
                earliest_new_date = row['date']
        _insert_returning_ids(Transaction, new_transactions)
        counts['transactions_created'] += len(new_transactions)

        # Only an already existing transaction can already carry this expense.
        reused_ids = [ref for row, ref in matched if row['expense'] and isinstance(ref, int)]
        existing_expenses = set()
        transactions_with_expenses = set()
        if reused_ids:
            linked = Expense.objects.filter(user_id=user_id, transaction_id__in=reused_ids).values_list(
                'transaction_id', 'date', 'amount', 'description', 'category_id'
            )
            for tx_pk, exp_date, amount, description, category_id in linked:
                existing_expenses.add((tx_pk, exp_date, amount, (description or '').strip(), category_id))
                transactions_with_expenses.add(tx_pk)

        new_expenses = []
        for row, ref in matched:
            expense = row['expense']
            if not expense:
                counts['deposit_rows'] += 1
                continue
            tx_pk = ref if isinstance(ref, int) else ref.pk
            if tx_pk is None:
                raise ImportAborted(f"Row {row['row']}: the transaction was not saved")
            amount = Decimal(expense['amount'])
            category_id = category_ids[expense['category']]
            if (tx_pk, row['date'], amount, row['description'], category_id) in existing_expenses:
                counts['expenses_existing'] += 1
                continue
            if tx_pk in transactions_with_expenses and len(warnings) < MAX_REPORTED_WARNINGS:
                warnings.append(
                    f"Row {row['row']}: transaction '{row['transaction_id']}' already had a different "
                    f"expense linked; a new expense was added to it."
                )
            new_expenses.append(Expense(
                user_id=user_id,
                title=expense['title'],
                amount=amount,
                description=row['description'],
                date=row['date'],
                category_id=category_id,
                transaction_id=tx_pk,
                is_cash=False,
                tax_category=expense['tax_category'],
                tax_percent_used=Decimal(expense['tax_percent']),
                tax_amount=Decimal(expense['tax_amount']) if expense['tax_amount'] is not None else None,
            ))
        Expense.objects.bulk_create(new_expenses, batch_size=CHUNK_SIZE)
        counts['expenses_created'] += len(new_expenses)

        if on_progress:
            on_progress(start + len(chunk))

    if earliest_new_date is not None:
        _recalculate_balances(user_id, account_id, earliest_new_date)

    # Every row must be accounted for exactly once, otherwise roll the whole job back.
    transactions_seen = counts['transactions_created'] + counts['transactions_existing']
    expenses_seen = counts['expenses_created'] + counts['expenses_existing'] + counts['deposit_rows']
    if transactions_seen != len(rows) or expenses_seen != len(rows):
        raise ImportAborted('Row count mismatch after import')

    return {**counts, 'total_rows': len(rows), 'categories_created': created_categories, 'warnings': warnings}


def process_import_job(job_id):
    """Entry point for the Celery task. Safe to call twice: only a pending job is ever picked up."""
    now = timezone.now()
    claimed = ExpenseImportJob.objects.filter(pk=job_id, status='pending').update(
        status='running', started_at=now, updated_at=now
    )
    if not claimed:
        return

    job = ExpenseImportJob.objects.get(pk=job_id)
    try:
        if not job.company_account_id:
            raise ImportAborted('The company account was deleted before the import started')
        with db_transaction.atomic():
            result = run_import(job, _ProgressReporter(job_id))
            # Finishing the job is part of the same transaction as the data: the rows and the
            # "completed" status are committed together or not at all.
            now = timezone.now()
            finished = ExpenseImportJob.objects.filter(pk=job_id, status='running').update(
                status='completed', result=result, processed_rows=job.total_rows,
                payload=None, finished_at=now, updated_at=now,
            )
            if not finished:
                raise ImportAborted('The import was cancelled')
    except BaseException as exc:
        # BaseException so a worker time limit or shutdown also releases the lock.
        logger.exception('Expense import job %s failed', job_id)
        reason = str(exc).strip() or exc.__class__.__name__
        now = timezone.now()
        ExpenseImportJob.objects.filter(pk=job_id, status='running').update(
            status='failed',
            error=f'Import failed and nothing was saved. {reason}'[:1000],
            payload=None, finished_at=now, updated_at=now,
        )
        if not isinstance(exc, Exception):
            raise
    finally:
        try:
            cache.delete(_progress_key(job_id))
        except Exception:
            pass


# --------------------------------------------------------------------------------------
# Job helpers used by the API
# --------------------------------------------------------------------------------------

def latest_job(user):
    """The user's most recent import. One SELECT, so it is cheap to poll."""
    jobs = ExpenseImportJob.objects.filter(user=user).defer('payload').order_by('-id')
    job = jobs.first()
    if job is None or job.status not in ExpenseImportJob.ACTIVE_STATUSES:
        return job

    # A job whose worker died would otherwise block the user forever. The worker is killed
    # after EXPENSE_IMPORT_TIME_LIMIT, so past the stale cutoff nothing can still be writing.
    now = timezone.now()
    if job.updated_at < now - timedelta(seconds=settings.EXPENSE_IMPORT_STALE_AFTER):
        ExpenseImportJob.objects.filter(pk=job.pk, status__in=ExpenseImportJob.ACTIVE_STATUSES).update(
            status='failed',
            error='The import did not finish in time and was cancelled. Nothing was saved. Please upload the file again.',
            payload=None, finished_at=now, updated_at=now,
        )
        job = jobs.first()
    return job


def serialize_job(job):
    if job is None:
        return None
    is_active = job.status in ExpenseImportJob.ACTIVE_STATUSES
    processed = job.processed_rows
    if job.status == 'running':
        try:
            processed = cache.get(_progress_key(job.id)) or 0
        except Exception:
            processed = 0
    percent = 100 if job.status == 'completed' else (
        int(processed * 100 / job.total_rows) if job.total_rows else 0
    )
    return {
        'id': job.id,
        'status': job.status,
        'is_active': is_active,
        'file_name': job.file_name,
        'company_account': job.company_account_id,
        'total_rows': job.total_rows,
        'processed_rows': processed,
        'percent': min(percent, 100),
        'result': job.result,
        'error': job.error,
        'created_at': job.created_at,
        'started_at': job.started_at,
        'finished_at': job.finished_at,
    }
