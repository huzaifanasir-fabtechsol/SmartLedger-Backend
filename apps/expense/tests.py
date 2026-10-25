from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, connection, transaction as db_transaction
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.account.models import User
from apps.expense import bulk_import
from apps.expense.models import Expense, ExpenseCategory, ExpenseImportJob
from apps.expense.serializers import ExpenseSerializer
from apps.revenue.models import CompanyAccount, Transaction

URL = '/api/expenses/bulk-import/'
HEADERS = ['date', 'transection id', 'withdraw amount', 'deposite amount', 'description',
           'notes', 'title', 'catagory', 'tax way', 'tax rate']


def withdraw_row(tx_id='TXN-1', amount=1100, day='2026-01-15', description='Printer paper',
                 category='Office', tax_way='Inclusive', tax_rate=10, title='Paper', notes='note'):
    return [day, tx_id, amount, None, description, notes, title, category, tax_way, tax_rate]


def deposit_row(tx_id='DEP-1', amount=5000, day='2026-01-16', description='Customer payment', notes='note'):
    return [day, tx_id, None, amount, description, notes, None, None, None, None]


def make_xlsx(rows, headers=HEADERS, name='import.xlsx', percent_cells=()):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in percent_cells:
        ws[cell].number_format = '0%'
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@override_settings(CACHES={'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
class BulkImportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='owner', email='owner@example.com', password='x')
        self.account = CompanyAccount.objects.create(
            user=self.user, bank_name='Bank', account_number='1', branch_code='1', account_holder='Owner'
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        # The worker is simulated: uploads only queue the job, tests run it explicitly.
        patcher = mock.patch('apps.expense.views.run_expense_import.apply_async')
        self.apply_async = patcher.start()
        self.addCleanup(patcher.stop)

    def upload(self, rows, account=None, **kwargs):
        return self.client.post(URL, {
            'file': make_xlsx(rows, **kwargs),
            'company_account': (account or self.account).id,
        }, format='multipart')

    def run_import(self, rows, **kwargs):
        response = self.upload(rows, **kwargs)
        self.assertEqual(response.status_code, 202, response.data)
        job_id = response.data['job']['id']
        bulk_import.process_import_job(job_id)
        return ExpenseImportJob.objects.get(pk=job_id)

    def assert_nothing_imported(self):
        self.assertEqual(Transaction.objects.count(), 0)
        self.assertEqual(Expense.objects.count(), 0)
        self.assertEqual(ExpenseCategory.objects.count(), 0)

    # ---- column / row validation ------------------------------------------------------

    def test_missing_columns_are_reported_and_nothing_is_queued(self):
        headers = [h for h in HEADERS if h not in ('tax rate', 'notes')]
        response = self.upload([], headers=headers)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Missing column(s): Notes, Tax Rate', response.data['error'])
        self.assertEqual(ExpenseImportJob.objects.count(), 0)

    def test_duplicate_column_is_rejected(self):
        response = self.upload([], headers=HEADERS + ['Date'])
        self.assertEqual(response.status_code, 400)
        self.assertIn('Duplicate column(s): Date', response.data['error'])

    def test_correctly_spelled_headers_in_any_order_are_accepted(self):
        headers = ['Tax Rate (%)', 'Tax Way', 'Category', 'Title', 'Notes', 'Description',
                   'Deposit Amount', 'Withdraw Amount', 'Transaction ID', 'Date']
        row = [10, 'inclusive', 'Office', 'Paper', 'note', 'Printer paper', None, 1100, 'TXN-1', '2026-01-15']
        job = self.run_import([row], headers=headers)
        self.assertEqual(job.status, 'completed')
        self.assertEqual(Expense.objects.get().title, 'Paper')

    def test_every_invalid_row_is_reported_and_nothing_is_saved(self):
        rows = [
            withdraw_row(),                                              # row 2: fine
            withdraw_row(tx_id=None),                                    # row 3
            withdraw_row(day='not a date'),                              # row 4
            [date(2026, 1, 15), 'T', None, None, 'd', 'n', 't', 'c', 'inclusive', 10],   # row 5: no amount
            [date(2026, 1, 15), 'T', 10, 10, 'd', 'n', 't', 'c', 'inclusive', 10],       # row 6: both amounts
            withdraw_row(description=None),                              # row 7
            withdraw_row(notes='  '),                                    # row 8
            withdraw_row(title=None),                                    # row 9
            withdraw_row(category=None),                                 # row 10
            withdraw_row(tax_way='maybe'),                               # row 11
            withdraw_row(tax_rate=None),                                 # row 12
            withdraw_row(tax_rate=150),                                  # row 13
            withdraw_row(tax_rate=0.1),                                  # row 14: fraction, not a percent
            withdraw_row(amount=-5),                                     # row 15
            withdraw_row(amount='abc'),                                  # row 16
            deposit_row(description=None),                               # row 17
        ]
        response = self.upload(rows)
        self.assertEqual(response.status_code, 400)
        errors = response.data['row_errors']
        self.assertEqual(response.data['total_errors'], 15)
        self.assertEqual([e.split(':')[0] for e in errors], [f'Row {n}' for n in range(3, 18)])
        expected = {
            3: 'Transaction ID is empty', 4: "Invalid Date 'not a date'",
            5: 'Either Withdraw Amount or Deposit Amount is required',
            6: 'Only one of Withdraw Amount or Deposit Amount can be filled',
            7: 'Description is empty', 8: 'Notes is empty', 9: 'Title is empty',
            10: 'Category is empty', 11: "Invalid Tax Way 'maybe'", 12: 'Tax Rate is empty',
            13: "Invalid Tax Rate '150'", 14: 'looks like a fraction',
            15: "Invalid Withdraw Amount '-5'", 16: "Invalid Withdraw Amount 'abc'",
            17: 'Description is empty',
        }
        for error in errors:
            row_number = int(error.split(':')[0].replace('Row ', ''))
            self.assertIn(expected[row_number], error)
        self.assertEqual(ExpenseImportJob.objects.count(), 0)
        self.apply_async.assert_not_called()
        self.assert_nothing_imported()

    def test_blank_rows_are_skipped_but_keep_real_row_numbers(self):
        rows = [withdraw_row(), [None] * 10, withdraw_row(tx_id='TXN-2', title=None)]
        response = self.upload(rows)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data['row_errors'], ['Row 4: Title is empty'])

    def test_request_validation(self):
        self.assertEqual(self.client.post(URL, {'file': make_xlsx([withdraw_row()])}, format='multipart').status_code, 400)
        self.assertEqual(self.client.post(URL, {'company_account': self.account.id}, format='multipart').status_code, 400)
        self.assertEqual(self.upload([withdraw_row()], name='import.csv').status_code, 400)
        self.assertEqual(self.upload([]).status_code, 400)   # header only

        other = User.objects.create_user(username='other', email='other@example.com', password='x')
        foreign_account = CompanyAccount.objects.create(
            user=other, bank_name='B', account_number='2', branch_code='2', account_holder='Other'
        )
        self.assertEqual(self.upload([withdraw_row()], account=foreign_account).status_code, 400)
        self.assertEqual(ExpenseImportJob.objects.count(), 0)

    def test_corrupt_file_is_rejected(self):
        bad = SimpleUploadedFile('import.xlsx', b'this is not a workbook')
        response = self.client.post(URL, {'file': bad, 'company_account': self.account.id}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Could not read the file', response.data['error'])

    # ---- importing ----------------------------------------------------------------------

    def test_import_creates_transactions_expenses_and_balances(self):
        Transaction.objects.create(
            user=self.user, company_account=self.account, date=date(2026, 1, 1), transaction_id='OPEN',
            withdraw=0, deposit=10000, balance=10000, description='Opening',
        )
        rows = [
            withdraw_row(tx_id='TXN-1', amount=1100, tax_way='INCLUSIVE', tax_rate=10),
            withdraw_row(tx_id='TXN-2', amount=1000, tax_way='exclusive', tax_rate='8%', description='Lunch', category='food'),
            withdraw_row(tx_id='TXN-3', amount=500, tax_way='Excl', tax_rate=0, description='Stamp'),
            deposit_row(tx_id='DEP-1', amount=5000),
        ]
        response = self.upload(rows)
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.data['job']['status'], 'pending')
        self.assertEqual(response.data['summary'], {
            'total_rows': 4, 'expense_rows': 3, 'deposit_rows': 1, 'new_categories': ['Office', 'food'],
        })
        self.apply_async.assert_called_once()
        self.assertEqual(Expense.objects.count(), 0)   # nothing written until the worker runs

        bulk_import.process_import_job(response.data['job']['id'])
        job = ExpenseImportJob.objects.get()
        self.assertEqual(job.status, 'completed')
        self.assertIsNone(job.payload)
        self.assertEqual(job.processed_rows, 4)
        self.assertEqual(job.result['transactions_created'], 4)
        self.assertEqual(job.result['expenses_created'], 3)
        self.assertEqual(job.result['deposit_rows'], 1)
        self.assertEqual(job.result['categories_created'], ['Office', 'food'])

        inclusive = Expense.objects.get(transaction__transaction_id='TXN-1')
        self.assertEqual((inclusive.amount, inclusive.tax_amount, inclusive.tax_category, inclusive.tax_percent_used),
                         (Decimal('1000.00'), Decimal('100.00'), 'inclusive', Decimal('10.00')))
        self.assertEqual((inclusive.title, inclusive.description, inclusive.date, inclusive.is_cash),
                         ('Paper', 'Printer paper', date(2026, 1, 15), False))
        self.assertEqual(inclusive.category.name, 'Office')
        self.assertEqual(inclusive.user, self.user)

        exclusive = Expense.objects.get(transaction__transaction_id='TXN-2')
        self.assertEqual((exclusive.amount, exclusive.tax_amount, exclusive.tax_category, exclusive.tax_percent_used),
                         (Decimal('1000.00'), Decimal('80.00'), 'exclusive', Decimal('8.00')))

        untaxed = Expense.objects.get(transaction__transaction_id='TXN-3')
        self.assertEqual((untaxed.amount, untaxed.tax_amount, untaxed.tax_category), (Decimal('500.00'), None, 'exclusive'))

        deposit = Transaction.objects.get(transaction_id='DEP-1')
        self.assertEqual((deposit.deposit, deposit.withdraw, deposit.notes), (Decimal('5000.00'), Decimal('0.00'), 'note'))
        self.assertEqual(deposit.expenses.count(), 0)   # a deposit never becomes an expense

        balances = list(Transaction.objects.order_by('date', 'id').values_list('transaction_id', 'balance'))
        self.assertEqual(balances, [
            ('OPEN', Decimal('10000.00')), ('TXN-1', Decimal('8900.00')), ('TXN-2', Decimal('7900.00')),
            ('TXN-3', Decimal('7400.00')), ('DEP-1', Decimal('12400.00')),
        ])

    def test_tax_numbers_match_a_manually_entered_expense(self):
        cases = [('1100', '10', 'inclusive'), ('999', '8', 'inclusive'), ('1234.56', '10', 'exclusive'),
                 ('333', '8', 'exclusive'), ('500', '0', 'inclusive')]
        for amount, percent, tax_category in cases:
            serializer = ExpenseSerializer(data={
                'title': 't', 'amount': amount, 'date': '2026-01-01',
                'tax_category': tax_category, 'tax_percent_used': percent,
            })
            self.assertTrue(serializer.is_valid(), serializer.errors)
            expected = (serializer.validated_data['amount'], serializer.validated_data.get('tax_amount'))
            self.assertEqual(bulk_import.split_tax(Decimal(amount), Decimal(percent), tax_category), expected)

    def test_percent_formatted_tax_rate_cell(self):
        job = self.run_import([withdraw_row(tax_rate=0.1)], percent_cells=['J2'])
        self.assertEqual(job.status, 'completed')
        self.assertEqual(Expense.objects.get().tax_percent_used, Decimal('10.00'))

    def test_existing_category_is_matched_ignoring_case_and_spacing(self):
        category = ExpenseCategory.objects.create(user=self.user, name='SUPPLY OFFICE  事務用品費')
        job = self.run_import([withdraw_row(category=' supply office 事務用品費 ')])
        self.assertEqual(job.result['categories_created'], [])
        self.assertEqual(ExpenseCategory.objects.count(), 1)
        self.assertEqual(Expense.objects.get().category, category)

    def test_category_name_owned_by_another_account_fails_validation(self):
        other = User.objects.create_user(username='other', email='other@example.com', password='x')
        ExpenseCategory.objects.create(user=other, name='Office')
        response = self.upload([withdraw_row(category='Office')])
        self.assertEqual(response.status_code, 400)
        self.assertIn("Category 'Office' (row 2) cannot be created", response.data['row_errors'][0])

    def test_importing_the_same_file_twice_creates_nothing_new(self):
        rows = [withdraw_row(tx_id='TXN-1'), withdraw_row(tx_id='TXN-2', amount=300), deposit_row()]
        self.run_import(rows)
        balances = list(Transaction.objects.order_by('date', 'id').values_list('balance', flat=True))

        second = self.run_import(rows)
        self.assertEqual(second.status, 'completed')
        self.assertEqual(second.result['transactions_created'], 0)
        self.assertEqual(second.result['transactions_existing'], 3)
        self.assertEqual(second.result['expenses_created'], 0)
        self.assertEqual(second.result['expenses_existing'], 2)
        self.assertEqual(second.result['warnings'], [])
        self.assertEqual((Transaction.objects.count(), Expense.objects.count()), (3, 2))
        self.assertEqual(list(Transaction.objects.order_by('date', 'id').values_list('balance', flat=True)), balances)

    def test_repeated_identical_bank_lines_are_matched_one_to_one(self):
        # Two transfer fees already recorded for the day; the file lists three.
        for _ in range(2):
            Transaction.objects.create(
                user=self.user, company_account=self.account, date=date(2026, 1, 15), transaction_id='FEE',
                withdraw=275, deposit=0, balance=0, description='Transfer fee',
            )
        fee = withdraw_row(tx_id='FEE', amount=275, description='Transfer fee', tax_rate=0)
        job = self.run_import([fee, fee, fee])
        self.assertEqual(job.result['transactions_existing'], 2)
        self.assertEqual(job.result['transactions_created'], 1)
        self.assertEqual(job.result['expenses_created'], 3)
        self.assertEqual(Transaction.objects.count(), 3)
        # Each transaction got exactly one expense.
        self.assertTrue(all(tx.expenses.count() == 1 for tx in Transaction.objects.all()))

    def test_same_id_and_date_with_a_different_amount_is_a_different_transaction(self):
        Transaction.objects.create(
            user=self.user, company_account=self.account, date=date(2026, 1, 15), transaction_id='CARD',
            withdraw=0, deposit=1000000, balance=1000000, description='Card',
        )
        job = self.run_import([deposit_row(tx_id='CARD', amount=2000000, day='2026-01-15')])
        self.assertEqual(job.result['transactions_created'], 1)
        self.assertEqual(Transaction.objects.count(), 2)

    def test_transactions_of_other_accounts_and_users_are_never_matched(self):
        other_account = CompanyAccount.objects.create(
            user=self.user, bank_name='B2', account_number='2', branch_code='2', account_holder='Owner'
        )
        Transaction.objects.create(
            user=self.user, company_account=other_account, date=date(2026, 1, 15), transaction_id='TXN-1',
            withdraw=1100, deposit=0, balance=0, description='Printer paper',
        )
        job = self.run_import([withdraw_row(tx_id='TXN-1', amount=1100)])
        self.assertEqual(job.result['transactions_created'], 1)
        self.assertEqual(Expense.objects.get().transaction.company_account, self.account)

    def test_warns_when_existing_transaction_already_has_a_different_expense(self):
        tx = Transaction.objects.create(
            user=self.user, company_account=self.account, date=date(2026, 1, 15), transaction_id='TXN-1',
            withdraw=1100, deposit=0, balance=0, description='Printer paper',
        )
        manual = Expense.objects.create(user=self.user, title='Manual', amount=1100, date=date(2026, 1, 15), transaction=tx)
        job = self.run_import([withdraw_row(tx_id='TXN-1', amount=1100)])
        self.assertEqual(job.result['expenses_created'], 1)
        self.assertEqual(job.result['warnings'], [
            f"Row 2: transaction 'TXN-1' already had expense #{manual.id} 'Manual' 1,100.00 "
            f"(no category, 2026-01-15). Added expense 'Paper' 1,000.00 + tax 100.00 (Office, 2026-01-15). "
            f"If they are the same charge, delete one of them."
        ])

    def test_review_lists_transactions_that_now_have_several_expenses(self):
        tx = Transaction.objects.create(
            user=self.user, company_account=self.account, date=date(2026, 1, 15), transaction_id='TXN-1',
            withdraw=1100, deposit=0, balance=0, description='Printer paper',
        )
        old = Expense.objects.create(user=self.user, title='ETC', amount=1100, date=date(2026, 1, 10), transaction=tx)
        # TXN-2 is reused but had no expense; TXN-3 is new. Neither needs review.
        Transaction.objects.create(
            user=self.user, company_account=self.account, date=date(2026, 1, 15), transaction_id='TXN-2',
            withdraw=300, deposit=0, balance=0, description='Lunch',
        )
        job = self.run_import([
            withdraw_row(tx_id='TXN-1', amount=1100),
            withdraw_row(tx_id='TXN-2', amount=300, description='Lunch'),
            withdraw_row(tx_id='TXN-3', amount=500, description='Stamp'),
        ])

        with self.assertNumQueries(2):
            bulk_import.review_transactions(job)

        response = self.client.get(URL + 'review/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['job_id'], job.id)
        [reviewed] = response.data['transactions']
        self.assertEqual((reviewed['id'], reviewed['transaction_id'], reviewed['withdraw']), (tx.id, 'TXN-1', Decimal('1100.00')))
        added = Expense.objects.get(transaction=tx, title='Paper')
        self.assertEqual(
            [(e['id'], e['title'], e['category'], e['total_amount'], e['from_this_import']) for e in reviewed['expenses']],
            [(old.id, 'ETC', None, Decimal('1100.00'), False),
             (added.id, 'Paper', 'Office', Decimal('1100.00'), True)],
        )

        # Deleting the duplicate through the normal expense API resolves the review item.
        self.assertEqual(self.client.delete(f'/api/expenses/{old.id}/').status_code, 204)
        self.assertEqual(self.client.get(URL + 'review/').data['transactions'], [])
        self.assertEqual(self.client.get(URL + f'review/?job={job.id}').data['transactions'], [])

    def test_review_is_scoped_to_the_user_and_finished_jobs(self):
        self.assertEqual(self.client.get(URL + 'review/').data, {'job_id': None, 'transactions': []})
        self.assertEqual(self.client.get(URL + 'review/?job=abc').status_code, 404)

        other = User.objects.create_user(username='other', email='other@example.com', password='x')
        foreign_job = ExpenseImportJob.objects.create(user=other, file_name='x.xlsx', status='completed')
        self.assertEqual(self.client.get(URL + f'review/?job={foreign_job.id}').status_code, 404)

        pending = self.upload([withdraw_row()]).data['job']
        self.assertEqual(self.client.get(URL + f"review/?job={pending['id']}").data['transactions'], [])

    def test_large_file_spanning_many_chunks_stays_query_efficient(self):
        rows = [withdraw_row(tx_id=f'TXN-{i}', amount=100 + i, category=f'Cat {i % 3}') for i in range(1200)]
        response = self.upload(rows)
        self.assertEqual(response.status_code, 202)
        # 2400 records are written in batches. The exact statement count depends on the database
        # (SQLite splits inserts at 999 variables), but it must stay far below one query per row.
        with CaptureQueriesContext(connection) as queries:
            bulk_import.process_import_job(response.data['job']['id'])
        self.assertLess(len(queries), 60)
        job = ExpenseImportJob.objects.get()
        self.assertEqual(job.status, 'completed')
        self.assertEqual((Transaction.objects.count(), Expense.objects.count()), (1200, 1200))
        self.assertEqual(Transaction.objects.order_by('date', 'id').last().balance,
                         -sum(Decimal(100 + i) for i in range(1200)))

    # ---- all-or-nothing -----------------------------------------------------------------

    def test_failure_midway_rolls_back_everything(self):
        rows = [withdraw_row(tx_id=f'TXN-{i}', amount=100 + i) for i in range(6)]
        response = self.upload(rows)
        real_bulk_create = Expense.objects.bulk_create
        calls = []

        def fail_on_second_chunk(objs, *args, **kwargs):
            calls.append(len(objs))
            if len(calls) == 2:
                raise RuntimeError('disk full')
            return real_bulk_create(objs, *args, **kwargs)

        with mock.patch.object(bulk_import, 'CHUNK_SIZE', 2), \
                mock.patch.object(Expense.objects, 'bulk_create', side_effect=fail_on_second_chunk), \
                self.assertLogs('apps.expense.bulk_import', level='ERROR'):
            bulk_import.process_import_job(response.data['job']['id'])

        self.assertEqual(calls, [2, 2])   # the first chunk really was written before the failure
        job = ExpenseImportJob.objects.get()
        self.assertEqual(job.status, 'failed')
        self.assertIn('nothing was saved', job.error)
        self.assertIn('disk full', job.error)
        self.assertIsNone(job.payload)
        self.assert_nothing_imported()

        # The lock is released, so the same file can be uploaded again and now succeeds.
        retry = self.run_import(rows)
        self.assertEqual(retry.status, 'completed')
        self.assertEqual((Transaction.objects.count(), Expense.objects.count()), (6, 6))

    def test_job_cancelled_while_running_is_rolled_back(self):
        response = self.upload([withdraw_row()])
        job_id = response.data['job']['id']

        def expire_job(job, on_progress=None):
            result = real_run_import(job, on_progress)
            ExpenseImportJob.objects.filter(pk=job_id).update(status='failed')   # e.g. released as stale
            return result

        real_run_import = bulk_import.run_import
        with mock.patch.object(bulk_import, 'run_import', side_effect=expire_job), \
                self.assertLogs('apps.expense.bulk_import', level='ERROR'):
            bulk_import.process_import_job(job_id)
        self.assert_nothing_imported()

    def test_task_only_ever_runs_a_pending_job_once(self):
        response = self.upload([withdraw_row(), deposit_row()])
        job_id = response.data['job']['id']
        bulk_import.process_import_job(job_id)
        bulk_import.process_import_job(job_id)   # duplicate delivery
        self.assertEqual((Transaction.objects.count(), Expense.objects.count()), (2, 1))
        self.assertEqual(ExpenseImportJob.objects.get().status, 'completed')

    # ---- one import at a time -----------------------------------------------------------

    def test_second_upload_is_refused_while_an_import_is_active(self):
        first = self.upload([withdraw_row()])
        self.assertEqual(first.status_code, 202)
        second = self.upload([withdraw_row(tx_id='TXN-9')])
        self.assertEqual(second.status_code, 409)
        self.assertEqual(second.data['job']['id'], first.data['job']['id'])
        self.assertEqual(ExpenseImportJob.objects.count(), 1)
        self.apply_async.assert_called_once()

        # Another user is not blocked.
        other = User.objects.create_user(username='other', email='other@example.com', password='x')
        account = CompanyAccount.objects.create(
            user=other, bank_name='B', account_number='2', branch_code='2', account_holder='Other'
        )
        other_client = APIClient()
        other_client.force_authenticate(other)
        response = other_client.post(
            URL, {'file': make_xlsx([withdraw_row(category='Other cat')]), 'company_account': account.id}, format='multipart'
        )
        self.assertEqual(response.status_code, 202)

    def test_database_refuses_two_active_jobs_for_one_user(self):
        ExpenseImportJob.objects.create(user=self.user, company_account=self.account, file_name='a.xlsx', status='running')
        with self.assertRaises(IntegrityError), db_transaction.atomic():
            ExpenseImportJob.objects.create(user=self.user, company_account=self.account, file_name='b.xlsx')
        # Finished jobs do not count.
        ExpenseImportJob.objects.update(status='completed')
        ExpenseImportJob.objects.create(user=self.user, company_account=self.account, file_name='c.xlsx')

    def test_stale_job_is_released(self):
        first = self.upload([withdraw_row()])
        job_id = first.data['job']['id']
        ExpenseImportJob.objects.filter(pk=job_id).update(updated_at=timezone.now() - timedelta(hours=1))

        status_response = self.client.get(URL)
        self.assertEqual(status_response.data['job']['status'], 'failed')
        self.assertIn('did not finish in time', status_response.data['job']['error'])

        # A late delivery of the dead job must not import anything.
        bulk_import.process_import_job(job_id)
        self.assert_nothing_imported()
        self.assertEqual(self.upload([withdraw_row()]).status_code, 202)

    def test_upload_fails_cleanly_when_the_worker_queue_is_down(self):
        self.apply_async.side_effect = ConnectionError('redis down')
        response = self.upload([withdraw_row()])
        self.assertEqual(response.status_code, 503)
        job = ExpenseImportJob.objects.get()
        self.assertEqual(job.status, 'failed')
        self.assertIsNone(job.payload)
        self.assert_nothing_imported()

        self.apply_async.side_effect = None
        self.assertEqual(self.upload([withdraw_row()]).status_code, 202)   # not left locked

    # ---- status / template --------------------------------------------------------------

    def test_status_endpoint(self):
        self.assertEqual(self.client.get(URL).data, {'job': None})

        response = self.upload([withdraw_row(), deposit_row()])
        job_id = response.data['job']['id']
        pending = self.client.get(URL).data['job']
        self.assertEqual((pending['status'], pending['is_active'], pending['percent'], pending['total_rows']),
                         ('pending', True, 0, 2))

        ExpenseImportJob.objects.filter(pk=job_id).update(status='running')
        bulk_import._ProgressReporter(job_id)(1)
        running = self.client.get(URL).data['job']
        self.assertEqual((running['status'], running['processed_rows'], running['percent']), ('running', 1, 50))

        ExpenseImportJob.objects.filter(pk=job_id).update(status='pending')
        bulk_import.process_import_job(job_id)
        done = self.client.get(URL).data['job']
        self.assertEqual((done['status'], done['is_active'], done['percent'], done['processed_rows']),
                         ('completed', False, 100, 2))
        self.assertEqual(done['result']['expenses_created'], 1)
        self.assertNotIn('payload', done)

        # Users only ever see their own job.
        other = User.objects.create_user(username='other', email='other@example.com', password='x')
        other_client = APIClient()
        other_client.force_authenticate(other)
        self.assertEqual(other_client.get(URL).data, {'job': None})

    def test_endpoints_require_authentication(self):
        anonymous = APIClient()
        self.assertEqual(anonymous.get(URL).status_code, 401)
        self.assertEqual(anonymous.post(URL, {}, format='multipart').status_code, 401)
        self.assertEqual(anonymous.get(URL + 'template/').status_code, 401)

    def test_template_download_is_itself_a_valid_import_file(self):
        response = self.client.get(URL + 'template/')
        self.assertEqual(response.status_code, 200)
        template = SimpleUploadedFile('template.xlsx', response.content)
        payload, summary = bulk_import.build_payload(template, self.user)
        self.assertEqual(summary['total_rows'], 3)
        self.assertEqual(summary['expense_rows'], 2)
        self.assertEqual(summary['deposit_rows'], 1)

    def test_multiple_taxes_crud_and_counts(self):
        category = ExpenseCategory.objects.create(user=self.user, name='General')
        # Test creating expense with 8% exclusive, 7% inclusive, and 10% exclusive
        data = {
            'title': 'Multi Tax Expense',
            'amount': '10700.00',
            'date': '2026-10-15',
            'category': category.id,
            'taxes': [
                {'name': 'Tax 8% Excl', 'rate': 8.0, 'category': 'exclusive'},
                {'name': 'Tax 7% Incl', 'rate': 7.0, 'category': 'inclusive'},
                {'name': 'Tax 10% Excl', 'rate': 10.0, 'category': 'exclusive'},
            ]
        }
        serializer = ExpenseSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        expense = serializer.save(user=self.user)

        # 10700 with 7% inclusive gives base: 10000.00, inclusive tax: 700.00
        # 8% exclusive on 10000 gives: 800.00
        # 10% exclusive on 10000 gives: 1000.00
        self.assertEqual(expense.amount, Decimal('10000.00'))
        self.assertEqual(expense.number_of_taxes, 3)
        self.assertEqual(expense.number_of_exclusive, 2)
        self.assertEqual(expense.number_of_inclusive_tax, 1)
        self.assertEqual(expense.exclusive_tax_amount, Decimal('1800.00'))
        self.assertEqual(expense.inclusive_tax_amount, Decimal('700.00'))
        self.assertEqual(expense.tax_amount, Decimal('2500.00'))
        self.assertEqual(expense.total_amount, Decimal('12500.00'))
        self.assertEqual(expense.tax_category, 'mixed')
        self.assertEqual(len(expense.taxes), 3)



class MixedTaxReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='mixed', email='mixed@x.com', password='pw')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        cat = ExpenseCategory.objects.create(user=self.user, name='General')
        def make(title, amount, taxes):
            s = ExpenseSerializer(data={'title': title, 'amount': amount, 'date': '2026-10-15',
                                        'category': cat.id, 'taxes': taxes})
            self.assertTrue(s.is_valid(), s.errors)
            s.save(user=self.user)
        make('mixed', '10700.00', [{'rate': 8, 'category': 'exclusive'}, {'rate': 7, 'category': 'inclusive'}])
        make('excl', '1000.00', [{'rate': 10, 'category': 'exclusive'}])
        make('incl', '1100.00', [{'rate': 10, 'category': 'inclusive'}])
        make('none', '500.00', [])

    def test_tax_stats_splits_mixed_from_inclusive_and_exclusive(self):
        d = self.client.get('/api/expenses/tax_stats/').json()
        self.assertEqual(d['mixed_count'], 1)
        self.assertEqual(d['mixed_base_amount'], 10000.0)
        self.assertEqual(d['mixed_excl_tax_amount'], 800.0)
        self.assertEqual(d['mixed_incl_tax_amount'], 700.0)
        self.assertEqual(d['total_mixed_amount'], 11500.0)
        self.assertEqual(d['inclusive_count'], 1)
        self.assertEqual(d['inclusive_tax_amount'], 100.0)
        self.assertEqual(d['exclusive_count'], 1)
        self.assertEqual(d['exclusive_tax_amount'], 100.0)
        self.assertEqual(d['total_tax_paid'], 1700.0)
        self.assertEqual(d['grand_total_amount'], 10000 + 1000 + 1000 + 500 + 1700)
        ten = next(b for b in d['tax_rate_breakdown'] if b['rate'] == 10.0)
        self.assertEqual(ten['count'], 2)
        self.assertEqual(ten['base_amount'], 2000.0)

    def test_filter_and_export_include_mixed(self):
        res = self.client.get('/api/expenses/?tax_category=mixed').json()
        self.assertEqual(res['count'], 1)
        x = self.client.get('/api/expenses/export_tax_xlsx/')
        self.assertEqual(x.status_code, 200)
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(x.content)).active
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
        mixed = next(r for r in rows if r[2] == 'mixed')
        self.assertEqual(mixed[5], 'Tax Mixed')
        self.assertEqual(mixed[10:15], [10000.0, 800.0, 700.0, 1500.0, 11500.0])


@override_settings(CACHES={'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}})
class DuplicateExpenseTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='dup', email='dup@x.com', password='pw')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.cat = ExpenseCategory.objects.create(user=self.user, name='General')

    def make(self, title='Fuel', amount='1100.00', taxes=None, **extra):
        data = {'title': title, 'amount': amount, 'date': '2026-10-15', 'category': self.cat.id,
                'taxes': taxes if taxes is not None else [{'rate': 10, 'category': 'inclusive'}], **extra}
        s = ExpenseSerializer(data=data)
        self.assertTrue(s.is_valid(), s.errors)
        return s.save(user=self.user)

    def test_stats_and_list_count_only_fully_identical_rows(self):
        a, b, c = self.make(), self.make(), self.make()
        self.make(description='different')
        self.make(taxes=[{'rate': 8, 'category': 'inclusive'}])
        Expense.objects.filter(pk=b.pk).update(is_highlighted=True)  # UI flag is not data
        other = User.objects.create_user(username='o', email='o@x.com', password='pw')
        Expense.objects.create(user=other, title='Fuel', amount=a.amount, date=a.date)

        stats = self.client.get('/api/expenses/stats/').json()
        self.assertEqual(stats['duplicate_count'], 2)
        self.assertEqual(stats['duplicate_group_count'], 1)
        d = self.client.get('/api/expenses/duplicates/').json()
        self.assertEqual(d['group_count'], 1)
        self.assertEqual([e['id'] for e in d['groups'][0]['expenses']], [a.id, b.id, c.id])

    def test_filters_apply_to_duplicates(self):
        self.make(), self.make()
        self.assertEqual(self.client.get('/api/expenses/stats/?month=9').json()['duplicate_count'], 0)
        self.assertEqual(self.client.get('/api/expenses/stats/?month=10').json()['duplicate_count'], 1)

    def test_delete_duplicate_refuses_last_copy(self):
        a, b = self.make(), self.make()
        res = self.client.delete(f'/api/expenses/{a.id}/delete_duplicate/')
        self.assertEqual(res.status_code, 204)
        res = self.client.delete(f'/api/expenses/{b.id}/delete_duplicate/')
        self.assertEqual(res.status_code, 409)
        self.assertTrue(Expense.objects.filter(pk=b.pk).exists())
        self.assertEqual(self.client.get('/api/expenses/duplicates/').json()['groups'], [])

    def test_delete_duplicate_rejects_unique_and_foreign_expenses(self):
        solo = self.make(title='Solo')
        self.assertEqual(self.client.delete(f'/api/expenses/{solo.id}/delete_duplicate/').status_code, 409)
        other = User.objects.create_user(username='o2', email='o2@x.com', password='pw')
        theirs = Expense.objects.create(user=other, title='T', amount=1, date=solo.date)
        Expense.objects.create(user=other, title='T', amount=1, date=solo.date)
        self.assertEqual(self.client.delete(f'/api/expenses/{theirs.id}/delete_duplicate/').status_code, 404)
