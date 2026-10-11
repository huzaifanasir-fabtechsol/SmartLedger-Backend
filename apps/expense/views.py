from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser
from django.conf import settings
from django.db import IntegrityError, transaction as db_transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase import pdfmetrics
from apps.expense.models import Expense, ExpenseCategory, ExpenseImportJob, Restaurant, SparePart
from apps.expense.serializers import ExpenseSerializer, ExpenseCategorySerializer, RestaurantSerializer, SparePartSerializer
from apps.expense.bulk_import import (
    COLUMNS, MAX_REPORTED_ERRORS, ImportFileError, build_payload, latest_job, serialize_job,
)
from apps.expense.tasks import run_expense_import
from apps.revenue.models import CompanyAccount, Transaction
from apps.revenue.serializers import TransactionSerializer
from apps.account.models import User
from project.pagination import CustomPageNumberPagination

class ExpenseCategoryViewSet(viewsets.ModelViewSet):
    serializer_class = ExpenseCategorySerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return ExpenseCategory.objects.filter(user=self.request.user)

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)



    @action(detail=False, methods=['get'])
    def all(self, request):
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

def _safe_parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    val_str = str(value).strip()
    if not val_str:
        return None
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%m/%d/%Y'):
        try:
            return datetime.strptime(val_str, fmt).date()
        except (ValueError, TypeError):
            continue
    return None

class ExpenseViewSet(viewsets.ModelViewSet):
    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Expense.objects.filter(user=self.request.user)

        category = self.request.query_params.get('category')
        date_val = self.request.query_params.get('date')
        date_from = self.request.query_params.get('date_from') or self.request.query_params.get('start_date')
        date_to = self.request.query_params.get('date_to') or self.request.query_params.get('end_date')
        year = self.request.query_params.get('year')
        month = self.request.query_params.get('month')
        search = self.request.query_params.get('search')

        if category:
            queryset = queryset.filter(category_id=category)
        if date_val:
            parsed_date = _safe_parse_date(date_val)
            if parsed_date:
                queryset = queryset.filter(date=parsed_date)
            else:
                queryset = queryset.none()
        if date_from:
            parsed_from = _safe_parse_date(date_from)
            if parsed_from:
                queryset = queryset.filter(date__gte=parsed_from)
        if date_to:
            parsed_to = _safe_parse_date(date_to)
            if parsed_to:
                queryset = queryset.filter(date__lte=parsed_to)
        if year:
            try:
                queryset = queryset.filter(date__year=int(year))
            except (ValueError, TypeError):
                pass
        if month:
            try:
                queryset = queryset.filter(date__month=int(month))
            except (ValueError, TypeError):
                pass
        payment_status = self.request.query_params.get('payment_status') or self.request.query_params.get('is_cash')
        if payment_status in ['true', 'True', '1', 'cash', 'Cash']:
            queryset = queryset.filter(is_cash=True)
        elif payment_status in ['false', 'False', '0', 'bank', 'Bank']:
            queryset = queryset.filter(is_cash=False)

        tax_cat = self.request.query_params.get('tax_category')
        if tax_cat:
            if tax_cat == 'none':
                queryset = queryset.filter(Q(tax_percent_used__isnull=True) | Q(tax_percent_used=0))
            elif tax_cat in ['inclusive', 'exclusive']:
                queryset = queryset.filter(tax_category=tax_cat)

        if search:
            queryset = queryset.filter(
                Q(title__icontains=search) |
                Q(description__icontains=search) |
                Q(category__name__icontains=search) |
                Q(restaurant__name__icontains=search) |
                Q(spare_part__name__icontains=search) |
                Q(spare_part__address__icontains=search)
            )

        return queryset.order_by('-date', '-id')

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)


    def _parse_excel_date(self, value):
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value

        value_str = str(value).strip()
        if not value_str:
            return None

        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%m/%d/%Y', '%Y%m%d'):
            try:
                return datetime.strptime(value_str, fmt).date()
            except ValueError:
                continue
        return None

    def _parse_excel_amount(self, value):
        if value is None:
            return None
        value_str = str(value).strip()
        if not value_str:
            return None

        cleaned = value_str.replace(',', '').replace('\u00a5', '').replace('$', '')
        cleaned = cleaned.replace('(', '-').replace(')', '')
        try:
            amount = Decimal(cleaned)
        except (InvalidOperation, TypeError):
            return None

        if amount < 0:
            amount = -amount
        return amount

    @action(detail=False, methods=['get'])
    def search_titles(self, request):
        query = request.query_params.get('q', '').strip()
        if not query:
            return Response([])
        
        titles = Expense.objects.filter(
            user=request.user,
            title__icontains=query
        ).values_list('title', flat=True).distinct()[:10]
        
        return Response(list(titles))

    def _add_watermark(self, canvas, doc, text="INVOICE"):
            canvas.saveState()
            canvas.setFont("HeiseiMin-W3", 80)
            canvas.setFillColor(colors.lightgrey)
            canvas.setFillAlpha(0.15)

            width, height = doc.pagesize
            canvas.translate(width / 2, height / 2)
            canvas.rotate(45)
            canvas.drawCentredString(0, 0, text)

            canvas.restoreState()

    def _add_page_decorations(self, canvas, doc, include_logo=True):
        from django.conf import settings
        import os
        from reportlab.lib.utils import ImageReader
        self._add_watermark(canvas, doc, "Ilyas Sons 合同会社")

        if include_logo:
            logo_path = os.path.join(settings.MEDIA_ROOT, "logo.png")
            print("MEDIA_ROOT:", settings.MEDIA_ROOT)
            print("Logo path:", logo_path)
            print("Exists:", os.path.exists(logo_path))
            if os.path.exists(logo_path):
                logo = ImageReader(logo_path)

                logo_width = 120
                logo_height = 40

                x = doc.leftMargin
                y = doc.pagesize[1] - logo_height - 10

                canvas.drawImage(
                    logo,
                    x,
                    y,
                    width=logo_width,
                    height=logo_height,
                    mask='auto'
                )

    def _add_first_page_decorations(self, canvas, doc):
        self._add_page_decorations(canvas, doc, include_logo=True)

    def _add_later_page_decorations(self, canvas, doc):
        self._add_page_decorations(canvas, doc, include_logo=False)

    @action(detail=True, methods=['get'])
    def generate_receipt(self, request, pk=None):
        expense = self.get_object()
        pdfmetrics.registerFont(UnicodeCIDFont('HeiseiMin-W3'))
        user = request.user

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="expense_{expense.id}.pdf"'

        doc = SimpleDocTemplate(response, pagesize=A4, topMargin=20, bottomMargin=20, leftMargin=20, rightMargin=20)
        elements = []
        styles = getSampleStyleSheet()

        # Title
        title_style = ParagraphStyle('Title', parent=styles['Normal'], fontSize=24, alignment=TA_CENTER, fontName='HeiseiMin-W3')
        elements.append(Paragraph("Expense Receipt", title_style))
        elements.append(Spacer(1, 15))

        # Company info on right
        header_data = [
            ["", "", user.company_name],
            ["", "", user.company_address],
            ["", "", f"TEL/FAX: {user.company_phone}"],
            ["", "", user.business_registration],
            ["", "", f"Date: {expense.date}"]
        ]
        header_table = Table(header_data, colWidths=[doc.width * 0.4, doc.width * 0.2, doc.width * 0.4])
        header_table.setStyle(TableStyle([
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ('FONTNAME', (2, 0), (2, -1), 'HeiseiMin-W3')
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 10))

        from reportlab.platypus import HRFlowable
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.black))
        elements.append(Spacer(1, 15))

        # Expense details with tax breakdown
        has_tax = expense.tax_percent_used is not None and expense.tax_percent_used > 0
        tax_amount = expense.tax_amount if expense.tax_amount is not None else (
            (expense.amount * expense.tax_percent_used / Decimal('100')) if has_tax else Decimal('0')
        )
        total_with_tax = expense.amount + (tax_amount if has_tax else Decimal('0'))

        tax_cat = getattr(expense, 'tax_category', 'exclusive')
        if has_tax:
            tax_cat_text = "Tax Inclusive" if tax_cat == 'inclusive' else "Tax Exclusive"
        else:
            tax_cat_text = "No Tax"

        detail_data = [
            ["Title:", expense.title],
            ["Base Amount:" if has_tax else "Amount:", f"¥ {expense.amount:,.0f}"],
            ["Tax Category:", tax_cat_text],
            ["Tax Rate:", f"{expense.tax_percent_used:.2f}%" if has_tax else "-"],
            ["Tax Amount:", f"¥ {tax_amount:,.0f}" if has_tax else "¥ 0"],
            ["Total (Incl. Tax):", f"¥ {total_with_tax:,.0f}"],
        ]

        detail_data.extend([
            ["Category:", expense.category.name if expense.category else 'N/A'],
            ["Description:", expense.description or '-']
        ])

        if expense.transaction:
            detail_data.extend([
                ["Transaction ID:", str(expense.transaction.transaction_id or expense.transaction.id)],
                ["Transaction:", expense.transaction.description],
                ["Transaction Amount:", f"¥ {expense.transaction.withdraw:,.0f}"]
            ])

        if expense.restaurant:
            detail_data.append(["Restaurant:", expense.restaurant.name])
        if expense.spare_part:
            spare_part_text = expense.spare_part.name
            if expense.spare_part.address:
                spare_part_text = f"{spare_part_text} - {expense.spare_part.address}"
            detail_data.append(["Shop / Vendor:", spare_part_text])

        if expense.is_cash:
            detail_data.append(["Payment Method:", "Cash"])
        elif expense.transaction:
            detail_data.append(["Payment Method:", "Bank / Account"])

        detail_table = Table(detail_data, colWidths=[doc.width * 0.3, doc.width * 0.7])
        detail_table.setStyle(TableStyle([
            ('FONTSIZE', (0, 0), (-1, -1), 11),
            ('FONTNAME', (0, 0), (-1, -1), 'HeiseiMin-W3'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8)
        ]))
        elements.append(detail_table)
        doc.build(
            elements,
            onFirstPage=self._add_first_page_decorations,
            onLaterPages=self._add_later_page_decorations,
        )
        return response

    @action(detail=False, methods=['get'])
    def export_pdf(self, request):
        pdfmetrics.registerFont(UnicodeCIDFont('HeiseiMin-W3'))
        user = request.user

        # Get filtered queryset
        queryset = self.get_queryset()
        date = request.query_params.get('date')
        date_from = request.query_params.get('date_from') or request.query_params.get('start_date')
        date_to = request.query_params.get('date_to') or request.query_params.get('end_date')
        year = request.query_params.get('year')
        month = request.query_params.get('month')
        category = request.query_params.get('category')

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="expenses.pdf"'

        doc = SimpleDocTemplate(response, pagesize=A4, topMargin=20, bottomMargin=20, leftMargin=20, rightMargin=20)
        elements = []
        styles = getSampleStyleSheet()

        # Title
        title_style = ParagraphStyle('Title', parent=styles['Normal'], fontSize=24, alignment=TA_CENTER, fontName='HeiseiMin-W3')
        elements.append(Paragraph("Expense Report", title_style))
        elements.append(Spacer(1, 15))

        # Header with company info and filters
        header_data = []
        left_col = []
        if date:
            left_col.append(f"Date: {date}")
        elif date_from or date_to:
            left_col.append(f"Period: {date_from or ''} ~ {date_to or ''}")
        elif year or month:
            parts = []
            if year:
                parts.append(f"{year}")
            if month:
                parts.append(f"Month {month}")
            left_col.append(f"Period: {' '.join(parts)}")

        if category:
            cat = ExpenseCategory.objects.filter(id=category).first()
            if cat:
                left_col.append(f"Category: {cat.name}")

        right_col = [user.company_name, user.company_address, f"TEL/FAX: {user.company_phone}", user.business_registration]

        max_rows = max(len(left_col), len(right_col))
        for i in range(max_rows):
            header_data.append([
                left_col[i] if i < len(left_col) else "",
                "",
                right_col[i] if i < len(right_col) else ""
            ])

        header_table = Table(header_data, colWidths=[doc.width * 0.4, doc.width * 0.2, doc.width * 0.4])
        header_table.setStyle(TableStyle([
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ('FONTNAME', (2, 0), (2, -1), 'HeiseiMin-W3')
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 10))

        from reportlab.platypus import HRFlowable
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.black))
        elements.append(Spacer(1, 15))

        # Table
        cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='HeiseiMin-W3',
            fontSize=7.5,
            leading=9.5,
            wordWrap='CJK',
        )
        right_cell_style = ParagraphStyle(
            'TableCellRight',
            parent=cell_style,
            alignment=TA_RIGHT,
        )
        center_cell_style = ParagraphStyle(
            'TableCellCenter',
            parent=cell_style,
            alignment=TA_CENTER,
        )

        table_data = [['Sr', 'Date', 'Title', 'Category', 'Payment/Tx', 'Shop/Vendor', 'Pre-Tax (¥)', 'Tax Type', 'Rate', 'Tax (¥)', 'Total (¥)']]
        total_base = Decimal('0')
        total_tax = Decimal('0')
        total_grand = Decimal('0')

        for idx, expense in enumerate(queryset, 1):
            place_text = '-'
            if expense.spare_part and expense.restaurant:
                place_text = f"{expense.restaurant.name} / {expense.spare_part.name}"
            elif expense.restaurant:
                place_text = expense.restaurant.name
            elif expense.spare_part:
                place_text = expense.spare_part.name
                if expense.spare_part.address:
                    place_text = f"{place_text} - {expense.spare_part.address}"

            if expense.is_cash:
                payment_text = 'Cash'
            elif expense.transaction:
                payment_text = str(expense.transaction.transaction_id or '-')
            else:
                payment_text = '-'

            has_tax = expense.tax_percent_used is not None and expense.tax_percent_used > 0
            if has_tax:
                tax_cat_label = "Inclusive" if getattr(expense, 'tax_category', 'exclusive') == 'inclusive' else "Exclusive"
                tax_rate_text = f"{expense.tax_percent_used:.1f}%"
            else:
                tax_cat_label = "None"
                tax_rate_text = "-"

            tax_amount_val = expense.tax_amount if expense.tax_amount is not None else (
                (expense.amount * expense.tax_percent_used / Decimal('100')) if has_tax else Decimal('0')
            )
            tax_amount_text = f"¥ {tax_amount_val:,.0f}" if (has_tax and tax_amount_val > 0) else "-"
            row_total = expense.amount + (tax_amount_val if has_tax else Decimal('0'))

            table_data.append([
                str(idx),
                str(expense.date),
                Paragraph(expense.title or '', cell_style),
                Paragraph(expense.category.name if expense.category else '-', cell_style),
                Paragraph(payment_text, cell_style),
                Paragraph(place_text, cell_style),
                Paragraph(f"¥ {expense.amount:,.0f}", right_cell_style),
                Paragraph(tax_cat_label, center_cell_style),
                Paragraph(tax_rate_text, center_cell_style),
                Paragraph(tax_amount_text, right_cell_style),
                Paragraph(f"¥ {row_total:,.0f}", right_cell_style)
            ])
            total_base += expense.amount
            if has_tax:
                total_tax += tax_amount_val
            total_grand += row_total

        table_data.append([
            '',
            '',
            '',
            '',
            '',
            Paragraph('Total:', cell_style),
            Paragraph(f"¥ {total_base:,.0f}", right_cell_style),
            '',
            '',
            Paragraph(f"¥ {total_tax:,.0f}", right_cell_style),
            Paragraph(f"¥ {total_grand:,.0f}", right_cell_style)
        ])

        table = Table(table_data, colWidths=[18, 48, 78, 48, 48, 65, 52, 44, 36, 50, 68])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.black),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'HeiseiMin-W3'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('FONTSIZE', (0, 1), (-1, -1), 7.5),
            ('FONTNAME', (0, 1), (-1, -1), 'HeiseiMin-W3'),
            ('ALIGN', (0, 1), (1, -1), 'CENTER'),
            ('ALIGN', (6, 1), (6, -1), 'RIGHT'),
            ('ALIGN', (7, 1), (8, -1), 'CENTER'),
            ('ALIGN', (9, 1), (10, -1), 'RIGHT'),
            ('VALIGN', (0, 1), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('BACKGROUND', (0, -1), (-1, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black)
        ]))
        elements.append(table)

        doc.build(
            elements,
            onFirstPage=self._add_first_page_decorations,
            onLaterPages=self._add_later_page_decorations,
        )
        return response

    @action(detail=False, methods=['get'])
    def export_xlsx(self, request):
        queryset = self.get_queryset()

        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        headers = [
            'Sr', 'Date', 'Title', 'Category', 'Cash', 'Transaction ID',
            'Restaurant', 'Shop / Spare Part', 'Description',
            'Amount (Excl. Tax) (¥)', 'Tax Category', 'Tax Rate (%)', 'Tax Amount (¥)', 'Total Incl. Tax (¥)'
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            if col_idx in [1, 2, 5, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [10, 13, 14]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        total_amount = Decimal('0')
        total_tax_amount = Decimal('0')
        total_grand_amount = Decimal('0')
        row_num = 2

        for idx, expense in enumerate(queryset, 1):
            spare_part_text = '-'
            if expense.spare_part:
                spare_part_text = expense.spare_part.name
                if expense.spare_part.address:
                    spare_part_text = f"{spare_part_text} - {expense.spare_part.address}"

            tx_text = '-'
            if expense.transaction:
                tx_text = expense.transaction.transaction_id or str(expense.transaction.id)

            category_name = expense.category.name if expense.category else '-'
            restaurant_name = expense.restaurant.name if expense.restaurant else '-'
            amount_val = float(expense.amount) if expense.amount is not None else 0.0
            cash_text = 'Yes' if expense.is_cash else 'No'

            has_tax = expense.tax_percent_used is not None and expense.tax_percent_used > 0
            tax_rate_val = float(expense.tax_percent_used) if has_tax else 0.0
            tax_amount_val = float(expense.tax_amount) if (expense.tax_amount is not None) else (
                round(amount_val * tax_rate_val / 100.0, 2) if has_tax else 0.0
            )
            total_with_tax = amount_val + tax_amount_val
            if has_tax:
                tax_cat_text = 'Tax Inclusive' if getattr(expense, 'tax_category', 'exclusive') == 'inclusive' else 'Tax Exclusive'
            else:
                tax_cat_text = 'No Tax'

            ws.append([
                idx,
                str(expense.date),
                expense.title or '',
                category_name,
                cash_text,
                tx_text,
                restaurant_name,
                spare_part_text,
                expense.description or '',
                amount_val,
                tax_cat_text,
                f"{tax_rate_val:.2f}%" if has_tax else "-",
                tax_amount_val,
                total_with_tax
            ])

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if col_idx in [1, 2, 5, 11, 12]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in [10, 13, 14]:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    c.number_format = '¥#,##0'
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            if expense.amount:
                total_amount += expense.amount
            if has_tax:
                total_tax_amount += Decimal(str(tax_amount_val))
            total_grand_amount += Decimal(str(total_with_tax))
            row_num += 1

        # Total row
        ws.append(['Total', '', '', '', '', '', '', '', '', float(total_amount), '', '', float(total_tax_amount), float(total_grand_amount)])
        total_row = row_num
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=total_row, column=col_idx)
            c.font = bold_font
            c.fill = total_fill
            c.border = thin_border
            if col_idx in [10, 13, 14]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '¥#,##0'
            elif col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")

        # Set row heights
        ws.row_dimensions[1].height = 26
        for r in range(2, total_row + 1):
            ws.row_dimensions[r].height = 20

        # Adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'expenses_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    @action(detail=False, methods=['get'])
    def tax_stats(self, request):
        queryset = self.get_queryset()

        total_original_amount = Decimal('0')
        total_tax_paid = Decimal('0')
        grand_total_amount = Decimal('0')

        total_inclusive_amount = Decimal('0')
        inclusive_base_amount = Decimal('0')
        inclusive_tax_amount = Decimal('0')
        inclusive_count = 0

        total_exclusive_amount = Decimal('0')
        exclusive_base_amount = Decimal('0')
        exclusive_tax_amount = Decimal('0')
        exclusive_count = 0

        total_no_tax_amount = Decimal('0')
        no_tax_count = 0

        rate_buckets = {}

        for exp in queryset:
            amt = exp.amount or Decimal('0')
            has_tax = exp.tax_percent_used is not None and exp.tax_percent_used > Decimal('0')
            if exp.tax_amount is not None:
                tax = exp.tax_amount
            elif has_tax:
                tax = (amt * exp.tax_percent_used / Decimal('100')).quantize(Decimal('0.01'))
            else:
                tax = Decimal('0')

            tot = amt + tax
            total_original_amount += amt
            total_tax_paid += tax
            grand_total_amount += tot

            tax_cat = getattr(exp, 'tax_category', 'exclusive') or 'exclusive'

            if not has_tax:
                total_no_tax_amount += amt
                no_tax_count += 1
                rate_key = 0.0
            else:
                rate_key = float(exp.tax_percent_used)
                if tax_cat == 'inclusive':
                    total_inclusive_amount += tot
                    inclusive_base_amount += amt
                    inclusive_tax_amount += tax
                    inclusive_count += 1
                else:
                    total_exclusive_amount += tot
                    exclusive_base_amount += amt
                    exclusive_tax_amount += tax
                    exclusive_count += 1

            if rate_key not in rate_buckets:
                rate_buckets[rate_key] = {
                    'rate': rate_key,
                    'base_amount': Decimal('0'),
                    'tax_amount': Decimal('0'),
                    'total_amount': Decimal('0'),
                    'count': 0
                }
            rate_buckets[rate_key]['base_amount'] += amt
            rate_buckets[rate_key]['tax_amount'] += tax
            rate_buckets[rate_key]['total_amount'] += tot
            rate_buckets[rate_key]['count'] += 1

        breakdown_list = []
        for rk in sorted(rate_buckets.keys(), reverse=True):
            b = rate_buckets[rk]
            breakdown_list.append({
                'rate': b['rate'],
                'base_amount': float(b['base_amount']),
                'tax_amount': float(b['tax_amount']),
                'total_amount': float(b['total_amount']),
                'count': b['count'],
            })

        return Response({
            'total_original_amount': float(total_original_amount),
            'total_tax_paid': float(total_tax_paid),
            'grand_total_amount': float(grand_total_amount),
            'total_inclusive_amount': float(total_inclusive_amount),
            'inclusive_base_amount': float(inclusive_base_amount),
            'inclusive_tax_amount': float(inclusive_tax_amount),
            'inclusive_count': inclusive_count,
            'total_exclusive_amount': float(total_exclusive_amount),
            'exclusive_base_amount': float(exclusive_base_amount),
            'exclusive_tax_amount': float(exclusive_tax_amount),
            'exclusive_count': exclusive_count,
            'total_no_tax_amount': float(total_no_tax_amount),
            'no_tax_count': no_tax_count,
            'expense_count': queryset.count(),
            'tax_rate_breakdown': breakdown_list,
        })

    @action(detail=False, methods=['get'])
    def export_tax_xlsx(self, request):
        queryset = self.get_queryset()

        wb = Workbook()
        ws = wb.active
        ws.title = "Tax Report Expenses"

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        # Title Block
        ws.append(["Expense Tax Report"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Records: {queryset.count()}"])
        ws.append([])

        headers = [
            'Sr', 'Date', 'Title', 'Category', 'Payment Method',
            'Tax Category', 'Tax Rate (%)', 'Pre-Tax Amount (¥)', 'Tax Amount (¥)', 'Total Amount (¥)'
        ]
        ws.append(headers)
        header_row = 4

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [8, 9, 10]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        total_amount = Decimal('0')
        total_tax_amount = Decimal('0')
        total_grand_amount = Decimal('0')
        row_num = header_row + 1

        for idx, expense in enumerate(queryset, 1):
            category_name = expense.category.name if expense.category else '-'
            amount_val = float(expense.amount) if expense.amount is not None else 0.0
            payment_method = 'Cash' if expense.is_cash else 'Bank'

            has_tax = expense.tax_percent_used is not None and expense.tax_percent_used > 0
            tax_rate_val = float(expense.tax_percent_used) if has_tax else 0.0
            tax_amount_val = float(expense.tax_amount) if (expense.tax_amount is not None) else (
                round(amount_val * tax_rate_val / 100.0, 2) if has_tax else 0.0
            )
            total_with_tax = amount_val + tax_amount_val

            if has_tax:
                tax_cat_text = 'Tax Inclusive' if getattr(expense, 'tax_category', 'exclusive') == 'inclusive' else 'Tax Exclusive'
            else:
                tax_cat_text = 'No Tax'

            ws.append([
                idx,
                str(expense.date),
                expense.title or '',
                category_name,
                payment_method,
                tax_cat_text,
                f"{tax_rate_val:.2f}%" if has_tax else "-",
                amount_val,
                tax_amount_val,
                total_with_tax
            ])

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if col_idx in [1, 2, 5, 6, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in [8, 9, 10]:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    c.number_format = '¥#,##0'
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            if expense.amount:
                total_amount += expense.amount
            if has_tax:
                total_tax_amount += Decimal(str(tax_amount_val))
            total_grand_amount += Decimal(str(total_with_tax))
            row_num += 1

        # Total row at bottom
        ws.append(['Total', '', '', '', '', '', '', float(total_amount), float(total_tax_amount), float(total_grand_amount)])
        total_row = row_num
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=total_row, column=col_idx)
            c.font = bold_font
            c.fill = total_fill
            c.border = thin_border
            if col_idx in [8, 9, 10]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '¥#,##0'
            elif col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[header_row].height = 26
        for r in range(header_row + 1, total_row + 1):
            ws.row_dimensions[r].height = 20

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'Tax_Report_Expenses_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    @action(detail=False, methods=['get'])
    def available_transactions(self, request):
        search = request.query_params.get('search', '')
        date_val = request.query_params.get('date', '')
        account_id = request.query_params.get('account_id', '')

        queryset = Transaction.objects.filter(user=request.user)

        if account_id:
            queryset = queryset.filter(company_account_id=account_id)
        if search:
            queryset = queryset.filter(
                Q(description__icontains=search) |
                Q(notes__icontains=search) |
                Q(balance__icontains=search) |
                Q(deposit__icontains=search) |
                Q(withdraw__icontains=search)
            )
        if date_val:
            parsed_date = _safe_parse_date(date_val)
            if parsed_date:
                queryset = queryset.filter(date=parsed_date)
            else:
                queryset = queryset.none()

        serializer = TransactionSerializer(queryset.order_by('-date', '-id'), many=True)
        return Response(serializer.data)


    @action(detail=False, methods=['post'], url_path='bulk-import-xls-expenses')
    def bulk_import_xls_expenses(self, request):
        excel_file = request.FILES.get('file')
        if not excel_file:
            return Response({'error': 'file is required'}, status=status.HTTP_400_BAD_REQUEST)

        category = ExpenseCategory.objects.filter(id=17, user=request.user).first()
        if not category:
            return Response({'error': 'Expense category 17 not found for this user'}, status=status.HTTP_400_BAD_REQUEST)

        account_id = request.data.get('company_account_id')
        company_account = None
        if account_id:
            company_account = CompanyAccount.objects.filter(id=account_id, user=request.user).first()
            if not company_account:
                return Response({'error': 'company_account_id is invalid'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            workbook = load_workbook(excel_file, data_only=True)
        except Exception as exc:
            return Response(
                {'error': f'Unable to read Excel file. Upload a valid .xlsx file. {str(exc)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) <= 1:
            return Response({'error': 'Excel file has no data rows'}, status=status.HTTP_400_BAD_REQUEST)

        default_account = company_account or CompanyAccount.objects.filter(user=request.user).order_by('id').first()
        missing_account_needed = False
        created_transactions = 0
        reused_transactions = 0
        created_expenses = 0
        skipped_rows = 0
        row_errors = []

        with db_transaction.atomic():
            for row_index, row in enumerate(rows[1:], start=2):
                row_values = list(row or [])
                if len(row_values) < 3:
                    skipped_rows += 1
                    row_errors.append(f'Row {row_index}: expected at least 3 columns')
                    continue

                tx_id_raw = row_values[-1]
                tx_id = str(tx_id_raw).strip() if tx_id_raw is not None else ''
                tx_id = tx_id[:500]
                if not tx_id:
                    skipped_rows += 1
                    row_errors.append(f'Row {row_index}: transaction id is empty')
                    continue

                parsed_date = self._parse_excel_date(row_values[0])
                if not parsed_date:
                    skipped_rows += 1
                    row_errors.append(f'Row {row_index}: invalid date "{row_values[0]}"')
                    continue

                amount = self._parse_excel_amount(row_values[1])
                if amount is None:
                    skipped_rows += 1
                    row_errors.append(f'Row {row_index}: invalid amount "{row_values[1]}"')
                    continue

                existing_tx = Transaction.objects.filter(
                    user=request.user,
                    transaction_id=tx_id
                ).order_by('id').first()

                if existing_tx:
                    tx = existing_tx
                    reused_transactions += 1
                else:
                    if not default_account:
                        missing_account_needed = True
                        skipped_rows += 1
                        row_errors.append(
                            f'Row {row_index}: transaction "{tx_id}" not found and no company account available to create it'
                        )
                        continue

                    tx = Transaction.objects.create(
                        user=request.user,
                        date=parsed_date,
                        transaction_id=tx_id,
                        withdraw=amount,
                        deposit=Decimal('0'),
                        balance=Decimal('0'),
                        description='Imported from expense xls',
                        notes='',
                        company_account=default_account,
                    )
                    created_transactions += 1

                Expense.objects.create(
                    user=request.user,
                    title='Highway',
                    category=category,
                    date=parsed_date,
                    amount=amount,
                    transaction=tx,
                    description='',
                )
                created_expenses += 1

        response_data = {
            'message': 'Bulk import completed',
            'created_transactions': created_transactions,
            'reused_transactions': reused_transactions,
            'created_expenses': created_expenses,
            'skipped_rows': skipped_rows,
            'errors': row_errors[:100],
        }
        if missing_account_needed:
            response_data['note'] = (
                'Provide company_account_id in the request or create a company account to allow creating new transactions.'
            )
        return Response(response_data, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get', 'post'], url_path='bulk-import', parser_classes=[MultiPartParser, FormParser])
    def bulk_import(self, request):
        """
        GET  -> status of the user's latest expense + transaction import.
        POST -> validate an .xlsx (file, company_account) and queue it for import.
        """
        if request.method == 'GET':
            return Response({'job': serialize_job(latest_job(request.user))})

        account_id = request.data.get('company_account') or request.data.get('company_account_id')
        if not account_id:
            return Response({'error': 'Company account is required.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            company_account = CompanyAccount.objects.filter(id=account_id, user=request.user).first()
        except (ValueError, TypeError):
            company_account = None
        if not company_account:
            return Response({'error': 'Selected company account not found or access denied.'}, status=status.HTTP_400_BAD_REQUEST)

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'error': 'No file uploaded. Please upload an .xlsx file.'}, status=status.HTTP_400_BAD_REQUEST)
        if not file_obj.name.lower().endswith('.xlsx'):
            return Response({'error': 'Invalid file format. Please upload an Excel .xlsx file.'}, status=status.HTTP_400_BAD_REQUEST)
        if file_obj.size > settings.EXPENSE_IMPORT_MAX_FILE_MB * 1024 * 1024:
            return Response(
                {'error': f'File is too large. The maximum size is {settings.EXPENSE_IMPORT_MAX_FILE_MB} MB.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        running_job = latest_job(request.user)
        if running_job and running_job.status in ExpenseImportJob.ACTIVE_STATUSES:
            return Response({
                'error': 'An import is already running. Please wait for it to finish.',
                'job': serialize_job(running_job),
            }, status=status.HTTP_409_CONFLICT)

        # Validate the whole file first. Nothing is saved unless every row is valid.
        try:
            payload, summary = build_payload(file_obj, request.user)
        except ImportFileError as exc:
            return Response({
                'error': exc.message,
                'total_errors': len(exc.row_errors),
                'row_errors': exc.row_errors[:MAX_REPORTED_ERRORS],
            }, status=status.HTTP_400_BAD_REQUEST)

        try:
            with db_transaction.atomic():
                job = ExpenseImportJob.objects.create(
                    user=request.user,
                    company_account=company_account,
                    file_name=file_obj.name[:255],
                    total_rows=summary['total_rows'],
                    payload=payload,
                )
        except IntegrityError:
            # Lost a race with another upload from the same user.
            return Response({
                'error': 'An import is already running. Please wait for it to finish.',
                'job': serialize_job(latest_job(request.user)),
            }, status=status.HTTP_409_CONFLICT)

        try:
            run_expense_import.apply_async(args=[job.id], expires=settings.EXPENSE_IMPORT_STALE_AFTER, retry=False)
        except Exception:
            ExpenseImportJob.objects.filter(pk=job.pk, status='pending').update(
                status='failed',
                error='The import could not be started because the background worker is unavailable. Nothing was saved.',
                payload=None,
                finished_at=timezone.now(),
            )
            return Response(
                {'error': 'The import could not be started right now. Nothing was saved. Please try again shortly.'},
                status=status.HTTP_503_SERVICE_UNAVAILABLE
            )

        job = ExpenseImportJob.objects.defer('payload').get(pk=job.pk)
        return Response({
            'message': 'File validated. The import is running in the background.',
            'job': serialize_job(job),
            'summary': summary,
        }, status=status.HTTP_202_ACCEPTED)

    @action(detail=False, methods=['get'], url_path='bulk-import/template')
    def bulk_import_template(self, request):
        category = ExpenseCategory.objects.filter(user=request.user).values_list('name', flat=True).first() or 'OFFICE SUPPLIES'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Expenses'
        ws.append([label for _, label, _ in COLUMNS])
        ws.append(['2026-01-15', 'TXN-0001', 1100, None, 'Printer paper', 'Paid by bank transfer', 'Office supplies', category, 'Inclusive', 10])
        ws.append(['2026-01-16', 'TXN-0002', None, 50000, 'Customer payment', 'Deposit rows do not create an expense', None, None, None, None])
        for idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=1, column=idx).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(idx)].width = 22

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="expense_import_template.xlsx"'
        wb.save(response)
        return response

class RestaurantViewSet(viewsets.ModelViewSet):
    serializer_class = RestaurantSerializer
    permission_classes = [IsAuthenticated]
    pagination_class = CustomPageNumberPagination

    def get_queryset(self):
        queryset = Restaurant.objects.filter(user=self.request.user)
        search = self.request.query_params.get('search', '')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) |
                Q(location__icontains=search)
            )
        return queryset

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)


class SparePartViewSet(viewsets.ModelViewSet):
    serializer_class = SparePartSerializer
    permission_classes = [IsAuthenticated]
    pagination_class = CustomPageNumberPagination

    def get_queryset(self):
        queryset = SparePart.objects.filter(user=self.request.user)
        search = self.request.query_params.get('search', '')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) |
                Q(address__icontains=search) |
                Q(description__icontains=search)
            )
        return queryset

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)
