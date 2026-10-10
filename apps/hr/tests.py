import datetime
from decimal import Decimal
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient
from rest_framework import status
import openpyxl

from apps.account.models import User
from apps.hr.models import Employee, Salary
from apps.hr.payroll_calculator import calculate_japanese_payroll
from apps.hr.payslip_generator import generate_payslip_excel, generate_payslip_pdf


class JapanesePayrollTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='testadmin',
            email='admin@example.com',
            password='Password123!',
            company_name='ILYAS SONS合同会社'
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        self.employee = Employee.objects.create(
            name='Hamza Muhammad',
            email='hamza@example.com',
            phone='090-1234-5678',
            role='Software Engineer',
            employment_start_month=datetime.date(2025, 1, 1),
            basic_salary=Decimal('200000'),
            commuting_allowance=Decimal('0'),
            dependents_count=0,
            employment_insurance_exempt=True,
            admin=self.user
        )

    def test_payroll_calculator_exact_match(self):
        """Test calculation produces exact figures from the Pay Slip.xlsx template."""
        res = calculate_japanese_payroll(
            basic_salary=Decimal('200000'),
            is_employment_insurance_exempt=True,
            dependents_count=0
        )
        self.assertEqual(res['basic_salary'], Decimal('200000'))
        self.assertEqual(res['taxable_payment'], Decimal('200000'))
        self.assertEqual(res['gross_payment'], Decimal('200000'))
        self.assertEqual(res['health_insurance'], Decimal('10030'))
        self.assertEqual(res['welfare_pension'], Decimal('18300'))
        self.assertEqual(res['employment_insurance'], Decimal('0'))
        self.assertEqual(res['total_social_insurance'], Decimal('28330'))
        self.assertEqual(res['taxable_income_base'], Decimal('171670'))
        self.assertEqual(res['income_tax'], Decimal('3770'))
        self.assertEqual(res['total_deductions'], Decimal('32100'))
        self.assertEqual(res['net_amount'], Decimal('167900'))

    def test_salary_create_auto_calculates(self):
        """Test creating salary via API automatically calculates taxes & deductions."""
        url = '/api/hr/salaries/'
        data = {
            'salary_month': '2026-02',
            'payment_date': '2026-02-28',
            'employee': self.employee.id,
            'working_days': 20,
            'working_hours': 160,
            'remarks': 'Monthly Salary'
        }
        resp = self.client.post(url, data, format='json')
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(float(resp.data['basic_salary']), 200000.0)
        self.assertEqual(float(resp.data['health_insurance']), 10030.0)
        self.assertEqual(float(resp.data['welfare_pension']), 18300.0)
        self.assertEqual(float(resp.data['total_social_insurance']), 28330.0)
        self.assertEqual(float(resp.data['income_tax']), 3770.0)
        self.assertEqual(float(resp.data['total_deductions']), 32100.0)
        self.assertEqual(float(resp.data['net_amount']), 167900.0)

    def test_calculate_preview_endpoint(self):
        """Test calculation preview endpoint without creating database record."""
        url = '/api/hr/salaries/calculate/'
        data = {
            'employee_id': self.employee.id,
            'overtime_allowance': 30000,
        }
        resp = self.client.post(url, data, format='json')
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(float(resp.data['taxable_payment']), 230000.0)
        self.assertGreater(float(resp.data['net_amount']), 0)

    def test_export_excel_endpoint(self):
        """Test exporting salary slip as Excel file."""
        salary = Salary.objects.create(
            salary_month='2026-02',
            payment_date=datetime.date(2026, 2, 28),
            employee=self.employee,
            admin=self.user,
            basic_salary=Decimal('200000'),
            taxable_payment=Decimal('200000'),
            gross_payment=Decimal('200000'),
            health_insurance=Decimal('10030'),
            welfare_pension=Decimal('18300'),
            employment_insurance=Decimal('0'),
            total_social_insurance=Decimal('28330'),
            taxable_income_base=Decimal('171670'),
            income_tax=Decimal('3770'),
            total_deductions=Decimal('32100'),
            net_amount=Decimal('167900'),
            working_days=Decimal('20'),
            working_hours=Decimal('160')
        )

        url = f'/api/hr/salaries/{salary.id}/export-excel/'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertTrue('Pay_Slip' in resp['Content-Disposition'])

        # Verify Excel content
        wb = openpyxl.load_workbook(filename=openpyxl.reader.excel.BytesIO(resp.content), data_only=False)
        ws = wb.active
        self.assertEqual(ws['C3'].value, 'Hamza Muhammad')
        self.assertEqual(ws['G13'].value, 200000)
        self.assertEqual(ws['R13'].value, 10030)
        self.assertEqual(ws['R14'].value, 18300)
        self.assertEqual(ws['R21'].value, 3770)

    def test_export_pdf_endpoint(self):
        """Test exporting salary slip as PDF."""
        salary = Salary.objects.create(
            salary_month='2026-02',
            payment_date=datetime.date(2026, 2, 28),
            employee=self.employee,
            admin=self.user,
            basic_salary=Decimal('200000'),
            taxable_payment=Decimal('200000'),
            gross_payment=Decimal('200000'),
            health_insurance=Decimal('10030'),
            welfare_pension=Decimal('18300'),
            employment_insurance=Decimal('0'),
            total_social_insurance=Decimal('28330'),
            taxable_income_base=Decimal('171670'),
            income_tax=Decimal('3770'),
            total_deductions=Decimal('32100'),
            net_amount=Decimal('167900')
        )

        url = f'/api/hr/salaries/{salary.id}/export-pdf/'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))
