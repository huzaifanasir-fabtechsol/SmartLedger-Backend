"""Excel version of the report package.

``XlsxWriter`` has the same interface as the PDF ``Writer``, so every document renderer produces
an Excel workbook with the same numbers: one sheet per form plus a contents sheet. Amounts are
written as real numbers so they can be summed and checked in Excel. Drawing-only elements
(the cover page, signature lines) have no Excel equivalent and are skipped.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .writer import Writer

AMOUNT_FORMAT = '#,##0;-#,##0'
GRAY = PatternFill('solid', fgColor='EDEDED')
THIN = Side(style='thin', color='000000')
MEDIUM = Side(style='medium', color='000000')
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED = Font(color='C00000', bold=True)
TEXT_COLS = 8                          # paragraphs/captions are merged across this many columns
POINTS_PER_CHAR = 5.5                  # PDF column widths (points) -> Excel character widths

_AMOUNT = re.compile(r'^-?\d{1,3}(,\d{3})*$')
_PERCENT = re.compile(r'^-?\d+(\.\d+)?%$')
_BAD_SHEET_CHARS = re.compile(r'[\[\]:*?/\\]')
_ALIGN = {'LEFT': 'left', 'RIGHT': 'right', 'CENTER': 'center', 'CENTRE': 'center'}


class _NullCanvas:
    """Absorbs direct canvas drawing (cover page, signature lines)."""

    def __getattr__(self, name):
        return lambda *args, **kwargs: None


def _cell_value(value):
    """'1,234' -> 1234 and '12.5%' -> 0.125 so Excel sees numbers; everything else stays text."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if _AMOUNT.match(text) and not (len(text) > 1 and text.lstrip('-').startswith('0')):
        return int(text.replace(',', ''))
    if _PERCENT.match(text):
        return float(text[:-1]) / 100
    return value


class XlsxWriter(Writer):
    def __init__(self, data, lang='ja', mode='full'):
        self.data, self.lang, self.mode = data, lang, mode
        self.font = None
        self.c = _NullCanvas()
        self.y = 0
        self.wb = Workbook()
        self.contents = self.wb.active
        self.contents.title = self.t(('目次', 'Contents'))
        self.ws = None
        self.row = 1
        self.sheets = []               # (sheet, title) in package order

    # ---- sheets -----------------------------------------------------------------------
    def _sheet_name(self, title):
        # The number prefix keeps names unique after Excel's 31-character cut
        return _BAD_SHEET_CHARS.sub(' ', f'{len(self.sheets) + 1:02d} {title}').strip()[:31]

    def begin_form(self, title, subtitle=None, filer=False, need=0):
        title_text = self.t(title)
        self.ws = self.wb.create_sheet(self._sheet_name(title_text))
        self.ws.sheet_view.showGridLines = False
        self.sheets.append((self.ws, title_text))
        self.row = 1
        self._text(title_text, font=Font(bold=True, size=14))
        if subtitle and self.t(subtitle):
            self._text(self.t(subtitle), font=Font(italic=True, color='555555'))
        self._draft_warning(self.ws)
        if filer and self.full:
            self.filer_block()
        elif filer:
            self._text(f"{self.company['name']}  /  {self.period_text()}")
        self.row += 1

    def _ensure_sheet(self):
        if self.ws is None:
            self.begin_form(('決算報告書・申告書類', 'Financial Statements & Tax Filing Package'))

    def _draft_warning(self, ws, row=None):
        bs = self.data.get('bs')
        if bs and not bs['balanced']:
            text = self.t((
                f"下書き: 貸借が一致していません (差額 {bs['difference']:,} 円)。提出しないでください。",
                f"DRAFT - balance sheet does not balance (difference {bs['difference']:,}). Do not file."))
            if row is None:
                self._text(text, font=RED)
            else:
                ws.cell(row=row, column=1, value=text).font = RED

    def new_page(self):
        """Page breaks have no meaning in a sheet; forms already get their own sheet."""

    def space(self, h=8):
        self._ensure_sheet()
        self.row += 1

    def _width(self, col, chars):
        letter = get_column_letter(col)
        current = self.ws.column_dimensions[letter].width or 0
        self.ws.column_dimensions[letter].width = max(current, min(chars, 70))

    # ---- text -------------------------------------------------------------------------
    def para(self, text, size=8, align=0):
        return self.t(text)

    def _text(self, text, font=None, wrap=False):
        ws = self.ws
        cell = ws.cell(row=self.row, column=1, value=text)
        if font:
            cell.font = font
        if wrap:
            ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=TEXT_COLS)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            lines = sum(len(part) // 110 + 1 for part in str(text).split('\n'))
            ws.row_dimensions[self.row].height = 15 * lines
        self.row += 1

    def caption(self, text, size=8, align='left'):
        self._ensure_sheet()
        self._text(self.t(text))

    def paragraph(self, text, size=8, gap=4):
        self._ensure_sheet()
        self._text(self.t(text), wrap=True)

    def lines(self, lines, size=9, leading=14, indent=0):
        self._ensure_sheet()
        for line in lines:
            self._text(self.t(line))

    def company_line(self, unit=True):
        self._ensure_sheet()
        self.ws.cell(row=self.row, column=1, value=self.company['name'])
        if unit:
            self.ws.cell(row=self.row, column=2, value=self.t(('（単位：　円）', '(Unit: JPY)'))).font = Font(color='555555')
        self.row += 1

    def filer_block(self):
        co = self.company
        self.kv_table([
            (('納税地', 'Address'), co['address']),
            (('電話', 'Phone'), co['phone']),
            (('法人名', 'Company name'), co['name']),
            (('法人番号', 'Corporate no.'), co['corp_no']),
            (('代表者', 'Representative'), co['rep']),
            (('税務署', 'Tax office'), co['tax_office']),
            (('事業年度', 'Fiscal period'), self.period_text()),
        ])

    # ---- tables -----------------------------------------------------------------------
    def table(self, data, widths, cmds=(), size=8, repeat=1, gap=10, x=None, keep=False):
        """Write a table given in reportlab form (rows of cells + TableStyle commands)."""
        self._ensure_sheet()
        if not data:
            return
        ws, top = self.ws, self.row
        n_rows, n_cols = len(data), max(len(r) for r in data)
        for r, row in enumerate(data):
            for c in range(n_cols):
                value = row[c] if c < len(row) else ''
                cell = ws.cell(row=top + r, column=c + 1, value=_cell_value(value) if value != '' else None)
                cell.border = GRID
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(cell.value, int):
                    cell.number_format = AMOUNT_FORMAT
                elif isinstance(cell.value, float):
                    cell.number_format = '0.0%'
        for c, width in enumerate(widths[:n_cols], start=1):
            self._width(c, width / POINTS_PER_CHAR)
        self._apply(cmds, top, n_rows, n_cols)
        self.row = top + n_rows + 1

    def _apply(self, cmds, top, n_rows, n_cols):
        def span(start, end):
            (c0, r0), (c1, r1) = start, end
            c0, c1 = (c if c >= 0 else n_cols + c for c in (c0, c1))
            r0, r1 = (r if r >= 0 else n_rows + r for r in (r0, r1))
            return range(top + r0, top + r1 + 1), range(c0 + 1, c1 + 2)

        ws = self.ws
        for cmd in cmds:
            name, start, end = cmd[0], cmd[1], cmd[2]
            rows, cols = span(start, end)
            if name == 'SPAN':
                if len(rows) > 1 or len(cols) > 1:
                    ws.merge_cells(start_row=rows[0], start_column=cols[0], end_row=rows[-1], end_column=cols[-1])
                continue
            for r in rows:
                for c in cols:
                    cell = ws.cell(row=r, column=c)
                    if name == 'BACKGROUND':
                        cell.fill = GRAY
                        cell.font = Font(bold=True)
                    elif name == 'ALIGN':
                        cell.alignment = Alignment(horizontal=_ALIGN.get(cmd[3], 'left'),
                                                   vertical='center', wrap_text=True)
                    elif name == 'LINEABOVE':
                        cell.border = Border(left=cell.border.left, right=cell.border.right,
                                             top=MEDIUM, bottom=cell.border.bottom)
                        cell.font = Font(bold=True)
                    elif name == 'LINEBELOW':
                        cell.border = Border(left=cell.border.left, right=cell.border.right,
                                             top=cell.border.top, bottom=MEDIUM)

    def kv_table(self, rows, widths=None, size=8):
        self._ensure_sheet()
        data = [[self.t(k), self.t(v)] for k, v in rows]
        self.table(data, widths or [150, 330], [('BACKGROUND', (0, 0), (0, -1), None)])

    def lines_table(self, columns, rows, size=8, no_header=False, extra_cmds=()):
        """Numbered-line form; values are written as numbers (not pre-formatted text)."""
        self._ensure_sheet()
        if not self.full:
            rows = self._simplify(rows)
        if not rows:
            return
        ws = self.ws
        n_cols = len(columns)
        if not no_header:
            for c, (header, _, _) in enumerate(columns, start=1):
                cell = ws.cell(row=self.row, column=c, value=self.t(header))
                cell.fill, cell.font, cell.border = GRAY, Font(bold=True), GRID
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            self.row += 1
        for row in rows:
            if row.kind == 'hdr':
                ws.cell(row=self.row, column=2, value=self.t(row.label))
                ws.merge_cells(start_row=self.row, start_column=2, end_row=self.row, end_column=n_cols)
                for c in range(1, n_cols + 1):
                    cell = ws.cell(row=self.row, column=c)
                    cell.fill, cell.font, cell.border = GRAY, Font(bold=True), GRID
                self.row += 1
                continue
            values = [str(row.no), self.t(row.label)] + list(row.vals)
            values += [None] * (n_cols - len(values))
            for c, value in enumerate(values[:n_cols], start=1):
                align = {'l': 'left', 'c': 'center', 'r': 'right'}[columns[c - 1][2]]
                cell = ws.cell(row=self.row, column=c, value=self._line_value(value))
                cell.border = GRID
                cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
                if isinstance(cell.value, int):
                    cell.number_format = AMOUNT_FORMAT
                if row.kind == 'total':
                    cell.font = Font(bold=True)
                    cell.border = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)
            self.row += 1
        for c, (_, width, _) in enumerate(columns, start=1):
            self._width(c, width / POINTS_PER_CHAR)
        self.row += 1

    def _line_value(self, value):
        if value is None or value == '':
            return None
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, float):
            return value
        if isinstance(value, int):
            return value
        if isinstance(value, (tuple, list)):
            return self.t(value)
        try:                                   # Decimal and other numerics -> whole yen
            return int(value)
        except (TypeError, ValueError):
            return _cell_value(str(value))

    # ---- output -----------------------------------------------------------------------
    def _write_contents(self):
        ws = self.contents
        ws.sheet_view.showGridLines = False
        co = self.company
        ws['A1'] = self.t(('決算報告書・申告書類', 'Financial Statements & Tax Filing Package'))
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = co['name']
        ws['A3'] = self.period_text()
        ws['A4'] = self.t(('参考資料(申告書そのものではありません)', 'Working copy - not a tax filing'))
        ws['A4'].font = Font(italic=True, color='555555')
        self._draft_warning(ws, row=5)
        row = 7
        ws.cell(row=row, column=1, value=self.t(('シート', 'Sheet'))).font = Font(bold=True)
        ws.cell(row=row, column=2, value=self.t(('書類', 'Document'))).font = Font(bold=True)
        for sheet, title in self.sheets:
            row += 1
            link = ws.cell(row=row, column=1, value=sheet.title)
            link.hyperlink = f"#'{sheet.title}'!A1"
            link.font = Font(color='0563C1', underline='single')
            ws.cell(row=row, column=2, value=title)
        ws.column_dimensions['A'].width = 34
        ws.column_dimensions['B'].width = 60

    def finish(self):
        self._write_contents()
        for sheet, _ in self.sheets:
            if (sheet.column_dimensions['A'].width or 0) < 12:
                sheet.column_dimensions['A'].width = 12
            sheet.page_setup.orientation = 'portrait'
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer
