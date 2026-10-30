import datetime
import io

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.account.models import User
from apps.expense.models import Expense, ExpenseCategory
from apps.reports.models import FiscalYear, FixedAsset, StatementInput
from apps.reports.pdf.package import PARTS, build_package
from apps.reports.services.depreciation import (declining_200_rate, depreciate, months_in_year,
                                                straight_line_rate)
from apps.reports.services.inputs import InputError, normalize_inputs
from apps.reports.services.package import compute_package
from apps.reports.services.statements import compute_statements
from apps.revenue.models import CompanyAccount, Order, Transaction

D = datetime.date

# Figures from the real FY2025 (第5期) 決算報告書 used as a golden fixture.
SGA_LINES = [
    ('役員報酬', 2400000), ('法定福利費', 116200), ('接待交際費', 268420), ('会議費', 392813),
    ('旅費交通費', 606261), ('通信費', 156816), ('消耗品費', 707399), ('事務用消耗品費', 19323),
    ('支払手数料', 781136), ('車両費', 2686506), ('燃料費', 800087), ('地代家賃', 1201547),
    ('保険料', 21200), ('租税公課', 232510), ('支払報酬料', 60000), ('雑費', 239696),
]
SAMPLE_INPUTS = {
    'sales': 28680691, 'purchases': 26895932, 'purchase_discounts': 11980,
    'closing_inventory': 2377190, 'cash': 84973,
    'other_current_assets': [
        {'label': '立替金', 'amount': 16820},
        {'label': '未収還付法人税等', 'amount': 554693},
        {'label': '役員貸付金', 'amount': 266560, 'executive': True},
    ],
    'investments': [{'label': '差入保証金', 'amount': 150000}],
    'current_liabilities': [
        {'label': '役員借入金', 'amount': 10229024}, {'label': '未払金', 'amount': 348494},
        {'label': '未払法人税等', 'amount': 204100}, {'label': '預り金', 'amount': 53240},
    ],
    'extra_sga': [{'label': l, 'amount': a} for l, a in SGA_LINES],
    'non_operating_income': [{'label': '受取利息', 'amount': 6}, {'label': '雑収入', 'amount': 594}],
    'income_taxes': 79100, 'prior_year_taxes': 150000,
    'capital': 5000000, 'opening_retained_earnings': -4064100,
    'tax': {
        'consumption': {
            'sales_incl_10': 31548751, 'nontaxable_sales': 6, 'purchase_incl_10': 35053275,
            'purchase_incl_8': 27857, 'purchase_returns_10': 13178,
            'transitional_incl_10': 3187439, 'transitional_incl_8': 56626,
        },
        'corporate': {'prior_losses': [{'label': '2023', 'amount': 136500}, {'label': '2024', 'amount': 327600}]},
    },
}


def make_sample_year(user):
    fy = FiscalYear.objects.create(user=user, term_no=5, start_date=D(2025, 1, 1), end_date=D(2025, 12, 31))
    StatementInput.objects.create(fiscal_year=fy, data=SAMPLE_INPUTS)
    assets = [
        ('簡易建物', 'building', D(2025, 9, 30), 728307, 'straight_line', 10),
        ('iPhone', 'equipment', D(2025, 2, 16), 108630, 'declining_200', 4),
        ('パソコン', 'equipment', D(2025, 5, 1), 119800, 'declining_200', 4),
        ('エアコン', 'equipment', D(2025, 4, 30), 112728, 'declining_200', 6),
    ]
    for name, cat, acquired, cost, method, life in assets:
        FixedAsset.objects.create(fiscal_year=fy, name=name, category=cat, acquired_on=acquired,
                                  cost=cost, method=method, useful_life=life)
    for name, acquired, cost in [('創立費', D(2021, 8, 10), 184700), ('開業費', D(2021, 12, 31), 321872)]:
        FixedAsset.objects.create(fiscal_year=fy, name=name, category='deferred', acquired_on=acquired,
                                  cost=cost, method='none', opening_book_value=cost)
    return fy


class DepreciationTests(TestCase):
    def test_statutory_rates(self):
        self.assertEqual(straight_line_rate(10), 0.100)
        self.assertEqual(straight_line_rate(3), 0.334)
        self.assertEqual(declining_200_rate(4), 0.500)
        self.assertEqual(declining_200_rate(6), 0.333)
        self.assertEqual(declining_200_rate(7), 0.286)

    def test_months_count_includes_acquisition_month(self):
        start, end = D(2025, 1, 1), D(2025, 12, 31)
        self.assertEqual(months_in_year(D(2025, 9, 30), start, end), 4)
        self.assertEqual(months_in_year(D(2025, 2, 16), start, end), 11)
        self.assertEqual(months_in_year(D(2023, 5, 1), start, end), 12)

    def test_matches_filed_depreciation(self):
        user = User.objects.create_user(username='u', email='u@example.com', password='x')
        fy = make_sample_year(user)
        got = {a.name: depreciate(a, fy.start_date, fy.end_date)['depreciation']
               for a in fy.fixed_assets.all()}
        self.assertEqual(got['簡易建物'], 24276)
        self.assertEqual(got['iPhone'], 49788)
        self.assertEqual(got['パソコン'], 39933)
        self.assertEqual(got['エアコン'], 28153)
        self.assertEqual(got['創立費'], 0)

    def test_lump_sum_writes_off_fully_and_floor_keeps_one_yen(self):
        user = User.objects.create_user(username='u', email='u@example.com', password='x')
        fy = FiscalYear.objects.create(user=user, term_no=1, start_date=D(2025, 1, 1), end_date=D(2025, 12, 31))
        lump = FixedAsset.objects.create(fiscal_year=fy, name='a', category='deferred', acquired_on=D(2025, 3, 1),
                                         cost=1000, method='lump_sum')
        nearly = FixedAsset.objects.create(fiscal_year=fy, name='b', category='equipment', acquired_on=D(2020, 3, 1),
                                           cost=100000, method='declining_200', useful_life=2, opening_book_value=1)
        self.assertEqual(depreciate(lump, fy.start_date, fy.end_date)['closing'], 0)
        self.assertEqual(depreciate(nearly, fy.start_date, fy.end_date)['closing'], 1)


class InputTests(TestCase):
    def test_defaults_and_rounding(self):
        out = normalize_inputs({'capital': 5000000.4})
        self.assertEqual(out['capital'], 5000000)
        self.assertIsNone(out['sales'])
        self.assertEqual(out['extra_sga'], [])

    def test_rejects_bad_types(self):
        for bad in ({'capital': 'abc'}, {'extra_sga': [{'amount': 5}]}, {'investments': 'x'}, 'nope'):
            with self.assertRaises(InputError):
                normalize_inputs(bad)


class StatementTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='u', email='u@example.com', password='x',
                                             company_name='ILYAS SONS 合同会社', company_address='愛知県大府市')
        self.fy = make_sample_year(self.user)

    def test_reproduces_filed_statements(self):
        data = compute_statements(self.user, self.fy)
        pl, bs, eq = data['pl'], data['bs'], data['equity']
        self.assertEqual(pl['cost_of_sales'], 24506762)
        self.assertEqual(pl['gross_profit'], 4173929)
        self.assertEqual(pl['sga_total'], 10832064)
        self.assertEqual(pl['operating_income'], -6658135)
        self.assertEqual(pl['ordinary_income'], -6657535)
        self.assertEqual(pl['taxes_total'], 229100)
        self.assertEqual(pl['net_income'], -6886635)
        self.assertEqual(bs['current_assets_total'], 3300236)
        self.assertEqual(bs['fixed_assets_total'], 1077315)
        self.assertEqual(bs['deferred_total'], 506572)
        self.assertEqual(bs['assets_total'], 4884123)
        self.assertEqual(bs['liabilities_total'], 10834858)
        self.assertEqual(bs['equity_total'], -5950735)
        self.assertTrue(bs['balanced'])
        self.assertEqual(eq['closing_retained'], -10950735)
        self.assertEqual(data['warnings'], [])

    def test_notes_include_derived_amounts(self):
        notes = compute_statements(self.user, self.fy)['notes']
        ja = '\n'.join(j for j, _ in notes)
        en = '\n'.join(e for _, e in notes)
        self.assertIn('266,560円', ja)
        self.assertIn('142,150円', ja)
        self.assertIn('JPY 266,560', en)
        self.assertIn('JPY 142,150', en)

    def test_unbalanced_sheet_is_flagged(self):
        record = self.fy.inputs
        record.data = {**SAMPLE_INPUTS, 'cash': 1}
        record.save()
        data = compute_statements(self.user, self.fy)
        self.assertFalse(data['bs']['balanced'])
        self.assertEqual(len(data['warnings']), 1)

    def test_derives_figures_from_app_data_when_not_overridden(self):
        record = self.fy.inputs
        record.data = {k: v for k, v in SAMPLE_INPUTS.items() if k not in ('sales', 'purchases', 'cash', 'extra_sga')}
        record.save()
        Order.objects.create(user=self.user, order_number='S1', transaction_type='sale', transaction_catagory='local',
                             transaction_date=D(2025, 6, 1), total_amount=1000000)
        Order.objects.create(user=self.user, order_number='S2', transaction_type='sale', transaction_catagory='local',
                             transaction_date=D(2024, 6, 1), total_amount=999)  # outside the year
        Order.objects.create(user=self.user, order_number='P1', transaction_type='purchase',
                             transaction_catagory='local', transaction_date=D(2025, 7, 1), total_amount=600000)
        category = ExpenseCategory.objects.create(name='燃料費', user=self.user)
        Expense.objects.create(user=self.user, title='gas', amount=5000, date=D(2025, 3, 3), category=category)
        account = CompanyAccount.objects.create(user=self.user, bank_name='b', account_number='1',
                                                branch_code='1', account_holder='h')
        for day, balance in ((D(2025, 5, 1), 111), (D(2025, 12, 30), 222), (D(2026, 1, 5), 999)):
            Transaction.objects.create(user=self.user, company_account=account, date=day, balance=balance,
                                       description='t')
        data = compute_statements(self.user, self.fy)
        self.assertEqual(data['pl']['sales'], 1000000)
        self.assertEqual(data['pl']['purchases'], 600000)
        self.assertEqual(data['bs']['current_assets'][0], {'label': '現金及び預金', 'amount': 222})
        self.assertIn({'label': '燃料費', 'amount': 5000}, data['pl']['sga_items'])
        self.assertEqual(data['sources']['sales'], 'derived')

    def test_other_users_data_is_excluded(self):
        other = User.objects.create_user(username='o', email='o@example.com', password='x')
        Order.objects.create(user=other, order_number='X1', transaction_type='sale', transaction_catagory='local',
                             transaction_date=D(2025, 6, 1), total_amount=777)
        record = self.fy.inputs
        record.data = {**SAMPLE_INPUTS, 'sales': None}
        record.save()
        self.assertEqual(compute_statements(self.user, self.fy)['pl']['sales'], 0)


class PdfTests(TestCase):
    def test_package_has_six_pages(self):
        user = User.objects.create_user(username='u', email='u@example.com', password='x', company_name='会社')
        data = compute_statements(user, make_sample_year(user))
        pdf = build_package(data).getvalue()
        self.assertTrue(pdf.startswith(b'%PDF'))
        self.assertEqual(pdf.count(b'/Type /Page\n') + pdf.count(b'/Type /Page '), 6)


class XlsxTests(TestCase):
    def test_every_part_renders_with_amounts_as_numbers(self):
        user = User.objects.create_user(username='u', email='u@example.com', password='x', company_name='会社')
        data = compute_package(user, make_sample_year(user))
        for lang in ('ja', 'en'):
            for mode in ('full', 'simple'):
                wb = load_workbook(build_package(data, list(PARTS), lang, mode, output='xlsx'))
                self.assertGreater(len(wb.sheetnames), 1)
        wb = load_workbook(build_package(data, ['kessan'], 'en', output='xlsx'))
        self.assertEqual(wb.sheetnames, ['Contents', '01 Balance Sheet', '02 Income Statement',
                                         '03 Schedule of SG&A Expenses', '04 Statement of Changes in Memb',
                                         '05 Notes to the Financial State'])
        bs = wb['01 Balance Sheet']
        total_row = [r for r in bs.iter_rows(values_only=True) if r[0] == 'Total assets'][0]
        self.assertEqual(total_row[1], data['bs']['assets_total'])
        self.assertEqual(total_row[3], data['bs']['liabilities_and_equity_total'])


class ApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='u', email='u@example.com', password='x', company_name='会社')
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_requires_authentication(self):
        resp = APIClient().get('/api/reports/fiscal-years/')
        self.assertIn(resp.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))

    def test_company_profile_prefilled_and_validated(self):
        resp = self.client.get('/api/reports/company-profile/')
        self.assertEqual(resp.data['legal_name'], '会社')
        bad = self.client.put('/api/reports/company-profile/', {'corporate_number': '123'}, format='json')
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)
        ok = self.client.put('/api/reports/company-profile/', {'corporate_number': '8120003020083'}, format='json')
        self.assertEqual(ok.status_code, status.HTTP_200_OK)

    def test_fiscal_year_crud_inputs_preview_and_pdf(self):
        resp = self.client.post('/api/reports/fiscal-years/',
                                {'term_no': 1, 'start_date': '2025-01-01', 'end_date': '2025-12-31'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        fy_id = resp.data['id']

        bad = self.client.put(f'/api/reports/fiscal-years/{fy_id}/inputs/', {'data': {'capital': 'x'}}, format='json')
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)
        ok = self.client.put(f'/api/reports/fiscal-years/{fy_id}/inputs/',
                             {'data': {'capital': 1000000, 'cash': 1000000}}, format='json')
        self.assertEqual(ok.status_code, status.HTTP_200_OK)
        self.assertEqual(self.client.get(f'/api/reports/fiscal-years/{fy_id}/inputs/').data['data']['capital'], 1000000)

        preview = self.client.get(f'/api/reports/fiscal-years/{fy_id}/preview/')
        self.assertEqual(preview.status_code, status.HTTP_200_OK)
        self.assertTrue(preview.data['bs']['balanced'])

        pdf = self.client.get(f'/api/reports/fiscal-years/{fy_id}/pdf/')
        self.assertEqual(pdf.status_code, status.HTTP_200_OK)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        unknown = self.client.get(f'/api/reports/fiscal-years/{fy_id}/pdf/?parts=nope')
        self.assertEqual(unknown.status_code, status.HTTP_400_BAD_REQUEST)

        xlsx = self.client.get(f'/api/reports/fiscal-years/{fy_id}/xlsx/?parts=kessan,hojinzei&lang=en&mode=simple')
        self.assertEqual(xlsx.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', xlsx['Content-Type'])
        self.assertIn('.xlsx', xlsx['Content-Disposition'])
        wb = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(wb.sheetnames[0], 'Contents')
        unknown = self.client.get(f'/api/reports/fiscal-years/{fy_id}/xlsx/?parts=nope')
        self.assertEqual(unknown.status_code, status.HTTP_400_BAD_REQUEST)

    def test_invalid_period_rejected(self):
        resp = self.client.post('/api/reports/fiscal-years/',
                                {'term_no': 1, 'start_date': '2025-12-31', 'end_date': '2025-01-01'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_cannot_access_other_users_year_or_attach_assets_to_it(self):
        other = User.objects.create_user(username='o', email='o@example.com', password='x')
        fy = FiscalYear.objects.create(user=other, term_no=1, start_date=D(2025, 1, 1), end_date=D(2025, 12, 31))
        self.assertEqual(self.client.get(f'/api/reports/fiscal-years/{fy.id}/preview/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/reports/fiscal-years/{fy.id}/pdf/').status_code, 404)
        resp = self.client.post('/api/reports/fixed-assets/', {
            'fiscal_year': fy.id, 'name': 'x', 'category': 'equipment', 'acquired_on': '2025-02-01',
            'cost': 1000, 'method': 'none'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_depreciating_asset_requires_useful_life(self):
        fy = FiscalYear.objects.create(user=self.user, term_no=1, start_date=D(2025, 1, 1), end_date=D(2025, 12, 31))
        resp = self.client.post('/api/reports/fixed-assets/', {
            'fiscal_year': fy.id, 'name': 'x', 'category': 'equipment', 'acquired_on': '2025-02-01',
            'cost': 1000, 'method': 'straight_line'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
