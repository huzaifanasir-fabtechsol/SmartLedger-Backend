"""
Letterhead Generator for Microsoft Word (.docx) export.
Produces an editable Word document matching company letterhead branding,
including full-bleed template background artwork, company logo/name typography,
clean blank writing canvas, and bottom date/contact footer.
"""

import io
import re
import base64
from datetime import datetime

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image, ImageDraw

# High resolution dimensions for A4 at 300 DPI (210 x 297 mm)
CANVAS_W = 2480
CANVAS_H = 3508
SCALE_X = CANVAS_W / 210.0
SCALE_Y = CANVAS_H / 297.0

# Template configuration matching LetterheadPrinter.jsx
TEMPLATE_CONFIGS = {
    'angles': {
        'top_in': 1.42,      # 36mm
        'bottom_in': 1.10,   # 28mm
        'is_centered': False,
    },
    'chevron': {
        'top_in': 1.02,      # 26mm
        'bottom_in': 1.10,   # 28mm
        'is_centered': False,
    },
    'wave': {
        'top_in': 1.26,      # 32mm
        'bottom_in': 1.10,   # 28mm
        'is_centered': True,
    },
    'corner': {
        'top_in': 0.71,      # 18mm
        'bottom_in': 1.10,   # 28mm
        'is_centered': False,
    },
    'classic': {
        'top_in': 0.79,      # 20mm
        'bottom_in': 1.10,   # 28mm
        'is_centered': True,
    },
}


def _clean_hex(hex_str: str, default: str = "1E3A8A") -> str:
    """Normalize hex color string to 6 uppercase characters without #."""
    if not hex_str:
        return default
    cleaned = re.sub(r'[^0-9a-fA-F]', '', str(hex_str)).upper()
    if len(cleaned) == 6:
        return cleaned
    if len(cleaned) == 3:
        return ''.join(c * 2 for c in cleaned)
    return default


def _hex_to_rgb(hex_str: str, default=(30, 58, 138)) -> RGBColor:
    """Convert hex string to docx RGBColor."""
    c = _clean_hex(hex_str)
    try:
        r = int(c[0:2], 16)
        g = int(c[2:4], 16)
        b = int(c[4:6], 16)
        return RGBColor(r, g, b)
    except Exception:
        return RGBColor(*default)


def _set_cell_no_borders(cell):
    """Remove borders from a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def render_template_artwork(template: str, c1_hex: str, c2_hex: str) -> io.BytesIO:
    """
    Render 300 DPI full-bleed artwork matching LetterheadPrinter.jsx canvas coordinates.
    The canvas is 210 x 297 mm.
    """
    c1 = f"#{c1_hex}"
    c2 = f"#{c2_hex}"

    img = Image.new('RGBA', (CANVAS_W, CANVAS_H), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)

    def pt(x, y):
        return (round(x * SCALE_X), round(y * SCALE_Y))

    def poly(pts, fill):
        draw.polygon([pt(x, y) for x, y in pts], fill=fill)

    def rect(x, y, w, h, fill):
        draw.rectangle([pt(x, y), pt(x + w, y + h)], fill=fill)

    tpl = template.lower() if template else 'angles'

    if tpl == 'angles':
        poly([(0, 0), (125, 0), (98, 17), (0, 17)], c1)
        poly([(0, 17), (98, 17), (86, 25), (0, 25)], c2)
        poly([(160, 0), (210, 0), (210, 14)], c2)
        poly([(210, 297), (210, 276), (150, 297)], c1)
        poly([(210, 297), (210, 287), (176, 297)], c2)
        poly([(0, 297), (0, 287), (30, 297)], c2)
    elif tpl == 'chevron':
        def chev_pts(x, y, w, h):
            return [
                (x, y),
                (x + w, y),
                (x + w + h / 2, y + h / 2),
                (x + w, y + h),
                (x, y + h),
                (x + h / 2, y + h / 2)
            ]
        for i, c in enumerate([c1, c2, c1]):
            poly(chev_pts(i * 10 - 4, 0, 7, 13), c)
            # 180 rotated at (210, 297)
            b_pts = [(210 - px, 297 - py) for px, py in chev_pts(i * 10 - 4, 0, 7, 13)]
            poly(b_pts, c)
        rect(0, 294, 150, 3, c1)
    elif tpl == 'wave':
        poly([(0, 0), (210, 0), (210, 10), (105, 22), (0, 10)], c1)
        poly([(0, 10), (105, 22), (210, 10), (210, 12.5), (105, 24.5), (0, 12.5)], c2)
        poly([(0, 297), (0, 291), (105, 284), (210, 291), (210, 297)], c1)
        poly([(0, 289), (105, 282), (210, 289), (210, 291), (105, 284), (0, 291)], c2)
    elif tpl == 'corner':
        poly([(210, 0), (210, 44), (118, 0)], c1)
        poly([(210, 0), (210, 28), (150, 0)], c2)
        poly([(0, 297), (0, 278), (42, 297)], c1)
        poly([(0, 297), (0, 288), (20, 297)], c2)
    else:  # classic
        rect(0, 0, 210, 7, c1)
        rect(0, 7, 210, 1.6, c2)
        rect(0, 291.2, 210, 1.2, c2)
        rect(0, 293, 210, 4, c1)

    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def _attach_background_artwork(section, artwork_buf: io.BytesIO):
    """
    Attach full-page background artwork to the section header anchored behind document text.
    Uses wp:anchor behindDoc="1" relativeFrom="page" spanning 8.27 x 11.69 inches.
    """
    header = section.header
    p_h = header.paragraphs[0]
    p_h.paragraph_format.space_before = Pt(0)
    p_h.paragraph_format.space_after = Pt(0)
    run_art = p_h.add_run()
    run_art.add_picture(artwork_buf, width=Inches(8.27), height=Inches(11.69))

    drawing = run_art._r.xpath('.//w:drawing')[0]
    inline = drawing.xpath('.//wp:inline')[0]
    graphic = inline.xpath('.//a:graphic')[0]

    # Full page extent in EMUs for 8.27 x 11.69 inches
    cx = 7562850
    cy = 10693400

    anchor_xml = f'''<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/>
  <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
  <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
  <wp:extent cx="{cx}" cy="{cy}"/>
  <wp:effectExtent l="0" t="0" r="0" b="0"/>
  <wp:wrapNone/>
</wp:anchor>'''
    anchor = parse_xml(anchor_xml)
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def generate_letterhead_docx(source) -> io.BytesIO:
    """
    Generate an editable Microsoft Word (.docx) document formatted with the
    company's letterhead design, branding, and contact details.

    Accepts either a Letterhead model instance or a dictionary of form values.
    """
    if hasattr(source, 'name'):
        company_name = source.name or 'Company Name'
        logo = source.logo or ''
        tagline = source.tagline or ''
        address = source.address or ''
        phone = source.phone or ''
        email = source.email or ''
        website = source.website or ''
        registration = source.registration or ''
        template = source.template or 'angles'
        color1 = source.color1 or '#1e3a8a'
        color2 = source.color2 or '#3b82f6'
        show_date = source.show_date
        date_str = datetime.now().strftime('%Y-%m-%d')
    elif isinstance(source, dict):
        company_name = source.get('name') or source.get('companyName') or 'Company Name'
        logo = source.get('logo', '')
        tagline = source.get('tagline', '')
        address = source.get('address', '')
        phone = source.get('phone', '')
        email = source.get('email', '')
        website = source.get('website', '')
        registration = source.get('registration', '')
        template = source.get('template', 'angles')
        color1 = source.get('color1', '#1e3a8a')
        color2 = source.get('color2', '#3b82f6')
        show_date = source.get('show_date', source.get('showDate', True))
        date_str = source.get('date') or datetime.now().strftime('%Y-%m-%d')
    else:
        company_name = 'Company Name'
        logo = ''
        tagline = ''
        address = ''
        phone = ''
        email = ''
        website = ''
        registration = ''
        template = 'angles'
        color1 = '#1e3a8a'
        color2 = '#3b82f6'
        show_date = True
        date_str = datetime.now().strftime('%Y-%m-%d')

    c1_hex = _clean_hex(color1, "1E3A8A")
    c2_hex = _clean_hex(color2, "3B82F6")
    c1_rgb = _hex_to_rgb(c1_hex)

    cfg = TEMPLATE_CONFIGS.get(template, TEMPLATE_CONFIGS['angles'])
    is_centered = cfg['is_centered']

    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(cfg['top_in'])
    section.bottom_margin = Inches(cfg['bottom_in'])
    section.left_margin = Inches(0.71)   # 18mm
    section.right_margin = Inches(0.71)  # 18mm
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0.65)

    # 1. Background Artwork (Full bleed, behind text)
    artwork_buf = render_template_artwork(template, c1_hex, c2_hex)
    _attach_background_artwork(section, artwork_buf)

    # 2. Decode Logo if present
    logo_bytes = None
    if logo and isinstance(logo, str) and 'base64,' in logo:
        try:
            _, b64_part = logo.split('base64,', 1)
            logo_bytes = base64.b64decode(b64_part)
        except Exception:
            logo_bytes = None

    # 3. Header Branding Section (Logo + Company Name + Tagline)
    if is_centered:
        if logo_bytes:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.paragraph_format.space_before = Pt(0)
            p_logo.paragraph_format.space_after = Pt(4)
            r_img = p_logo.add_run()
            r_img.add_picture(io.BytesIO(logo_bytes), height=Inches(0.8))

        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_name.paragraph_format.space_before = Pt(0)
        p_name.paragraph_format.space_after = Pt(2)
        r_name = p_name.add_run(company_name)
        r_name.font.name = 'Segoe UI'
        r_name.font.size = Pt(22)
        r_name.font.bold = True
        r_name.font.color.rgb = c1_rgb

        if tagline:
            p_tag = doc.add_paragraph()
            p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_tag.paragraph_format.space_before = Pt(0)
            p_tag.paragraph_format.space_after = Pt(20)
            r_tag = p_tag.add_run(tagline.upper())
            r_tag.font.name = 'Segoe UI'
            r_tag.font.size = Pt(8.5)
            r_tag.font.color.rgb = RGBColor(107, 114, 128)
        else:
            p_name.paragraph_format.space_after = Pt(20)
    else:
        # Left-Aligned Layout
        if logo_bytes:
            head_table = doc.add_table(rows=1, cols=2)
            head_table.alignment = WD_TABLE_ALIGNMENT.LEFT
            head_table.autofit = False

            c_logo = head_table.cell(0, 0)
            c_logo.width = Inches(1.8)
            _set_cell_no_borders(c_logo)
            c_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_l = c_logo.paragraphs[0]
            p_l.paragraph_format.space_before = Pt(0)
            p_l.paragraph_format.space_after = Pt(0)
            r_img = p_l.add_run()
            r_img.add_picture(io.BytesIO(logo_bytes), height=Inches(0.85))

            c_text = head_table.cell(0, 1)
            c_text.width = Inches(5.0)
            _set_cell_no_borders(c_text)
            c_text.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p_n = c_text.paragraphs[0]
            p_n.paragraph_format.space_before = Pt(0)
            p_n.paragraph_format.space_after = Pt(2)
            r_name = p_n.add_run(company_name)
            r_name.font.name = 'Segoe UI'
            r_name.font.size = Pt(22)
            r_name.font.bold = True
            r_name.font.color.rgb = c1_rgb

            if tagline:
                p_t = c_text.add_paragraph()
                p_t.paragraph_format.space_before = Pt(0)
                p_t.paragraph_format.space_after = Pt(0)
                r_tag = p_t.add_run(tagline.upper())
                r_tag.font.name = 'Segoe UI'
                r_tag.font.size = Pt(8.5)
                r_tag.font.color.rgb = RGBColor(107, 114, 128)
        else:
            p_name = doc.add_paragraph()
            p_name.paragraph_format.space_before = Pt(0)
            p_name.paragraph_format.space_after = Pt(2)
            r_name = p_name.add_run(company_name)
            r_name.font.name = 'Segoe UI'
            r_name.font.size = Pt(22)
            r_name.font.bold = True
            r_name.font.color.rgb = c1_rgb

            if tagline:
                p_tag = doc.add_paragraph()
                p_tag.paragraph_format.space_before = Pt(0)
                p_tag.paragraph_format.space_after = Pt(20)
                r_tag = p_tag.add_run(tagline.upper())
                r_tag.font.name = 'Segoe UI'
                r_tag.font.size = Pt(8.5)
                r_tag.font.color.rgb = RGBColor(107, 114, 128)
            else:
                p_name.paragraph_format.space_after = Pt(20)

    # 4. Clean, Blank Editable Body Canvas
    # Only a blank paragraph with standard body typography ready for user writing
    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_before = Pt(14)
    p_body.paragraph_format.space_after = Pt(6)
    p_body.paragraph_format.line_spacing = 1.15
    r_body = p_body.add_run()
    r_body.font.name = 'Segoe UI'
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = RGBColor(31, 41, 55)

    # 5. Section Footer: Date + Horizontal Dividing Rule + Contact Details
    footer = section.footer
    p_first = footer.paragraphs[0]

    contact_parts = []
    if address:
        contact_parts.append(("Address", address))
    if phone:
        contact_parts.append(("Phone", phone))
    if email:
        contact_parts.append(("Email", email))
    if website:
        contact_parts.append(("Web", website))
    if registration:
        contact_parts.append(("Reg. No", registration))

    # Date line
    if show_date and date_str:
        p_date = p_first
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_centered else WD_ALIGN_PARAGRAPH.LEFT
        p_date.paragraph_format.space_before = Pt(0)
        p_date.paragraph_format.space_after = Pt(4)

        r_dlbl = p_date.add_run("Date: ")
        r_dlbl.font.name = 'Segoe UI'
        r_dlbl.font.size = Pt(9.5)
        r_dlbl.font.bold = True
        r_dlbl.font.color.rgb = c1_rgb

        r_dval = p_date.add_run(str(date_str))
        r_dval.font.name = 'Segoe UI'
        r_dval.font.size = Pt(9.5)
        r_dval.font.color.rgb = RGBColor(55, 65, 81)

        p_contact = footer.add_paragraph() if contact_parts else None
    else:
        p_contact = p_first if contact_parts else None

    # Contact paragraph with top border line
    if p_contact and contact_parts:
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contact.paragraph_format.space_before = Pt(4)
        p_contact.paragraph_format.space_after = Pt(0)
        p_contact.paragraph_format.line_spacing = 1.25

        # Thin dividing rule matching border-top: 1.5px solid var(--lh-c1)
        pPr = p_contact._p.get_or_add_pPr()
        pBdr_xml = f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="6" w:color="{c1_hex}"/></w:pBdr>'
        pPr.append(parse_xml(pBdr_xml))

        for idx, (label, val) in enumerate(contact_parts):
            if idx > 0:
                if idx == 3 and len(contact_parts) > 3:
                    # Break into second line for clean visual balance matching Screenshot 1
                    p_contact.add_run('\n')
                else:
                    r_sp = p_contact.add_run('    ')
                    r_sp.font.size = Pt(8)

            r_l = p_contact.add_run(f"{label}: ")
            r_l.font.name = 'Segoe UI'
            r_l.font.size = Pt(8)
            r_l.font.bold = True
            r_l.font.color.rgb = c1_rgb

            r_v = p_contact.add_run(str(val))
            r_v.font.name = 'Segoe UI'
            r_v.font.size = Pt(8)
            r_v.font.color.rgb = RGBColor(75, 85, 99)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
