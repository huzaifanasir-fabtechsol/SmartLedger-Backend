import hashlib
import os
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APITestCase

from apps.account.models import Letterhead, LetterheadRecord, User


class LetterheadApiTests(APITestCase):
    url = '/api/account/letterheads/'

    def setUp(self):
        self.user = User.objects.create_user(username='a', email='a@example.com', password='pw-12345!')
        self.other = User.objects.create_user(username='b', email='b@example.com', password='pw-12345!')
        self.client.force_authenticate(self.user)

    def test_requires_login(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(self.url).status_code, 401)

    def test_create_update_delete(self):
        res = self.client.post(self.url, {'name': ' Safeer Motors ', 'template': 'wave', 'color1': '#112233'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.data['name'], 'Safeer Motors')
        item = f"{self.url}{res.data['id']}/"
        res = self.client.patch(item, {'color2': '#abcdef', 'tagline': 'Group'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(Letterhead.objects.get().color2, '#abcdef')
        self.assertEqual(self.client.delete(item).status_code, 204)
        self.assertEqual(Letterhead.objects.count(), 0)

    def test_only_own_records_visible(self):
        Letterhead.objects.create(user=self.other, name='Theirs')
        mine = Letterhead.objects.create(user=self.user, name='Mine')
        res = self.client.get(self.url)
        results = res.data['results'] if isinstance(res.data, dict) else res.data
        self.assertEqual([r['id'] for r in results], [mine.id])
        theirs = Letterhead.objects.get(name='Theirs')
        self.assertEqual(self.client.delete(f'{self.url}{theirs.id}/').status_code, 404)

    def test_validation(self):
        bad = [
            {'name': '   '},
            {'name': 'X', 'color1': 'red'},
            {'name': 'X', 'template': 'nope'},
            {'name': 'X', 'logo': 'http://evil/x.png'},
        ]
        for payload in bad:
            self.assertEqual(self.client.post(self.url, payload, format='json').status_code, 400, payload)

    def test_export_docx_from_payload(self):
        """Test exporting letterhead as editable docx directly from JSON payload."""
        import docx
        import io

        payload = {
            'companyName': 'ILYAS SONS 合同会社',
            'tagline': 'Automotive Group',
            'address': 'Tokyo, Japan',
            'phone': '+81 90-1234-5678',
            'email': 'contact@ilyassons.jp',
            'template': 'angles',
            'color1': '#1e3a8a',
            'color2': '#3b82f6',
            'showDate': True,
        }
        res = self.client.post(f"{self.url}export-docx/", payload, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        self.assertIn('.docx', res['Content-Disposition'])

        # Verify it is a valid Word document
        doc = docx.Document(io.BytesIO(res.content))
        self.assertTrue(len(doc.paragraphs) > 0)
        found_name = any('ILYAS SONS' in p.text for p in doc.paragraphs)
        self.assertTrue(found_name)

    def test_export_docx_from_saved(self):
        """Test exporting saved letterhead by ID as editable docx."""
        lh = Letterhead.objects.create(
            user=self.user,
            name='My Saved Letterhead',
            tagline='Corporate Group',
            address='Osaka, Japan',
            color1='#0f172a',
            color2='#10b981'
        )
        res = self.client.get(f"{self.url}{lh.id}/export-docx/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        self.assertIn('My_Saved_Letterhead', res['Content-Disposition'])


from django.test import TestCase as _LockTestCase
from rest_framework.test import APIClient as _LockClient


class LockPinTests(_LockTestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='u', email='u@x.com', password='Str0ng-Pass!')
        self.client = _LockClient()
        self.client.force_authenticate(self.user)

    def test_set_verify_remove(self):
        url = '/api/account/lock-pin/'
        self.assertEqual(self.client.post(url, {'pin': '123456'}, format='json').status_code, 400)
        self.assertEqual(self.client.put(url, {'pin': '123456', 'password': 'wrong'}, format='json').status_code, 400)
        self.assertEqual(self.client.put(url, {'pin': '12345', 'password': 'Str0ng-Pass!'}, format='json').status_code, 400)
        self.assertEqual(self.client.put(url, {'pin': '123456', 'password': 'Str0ng-Pass!'}, format='json').status_code, 200)
        self.assertTrue(self.client.get('/api/account/profile/').json()['has_lock_pin'])
        self.assertEqual(self.client.post(url, {'pin': '000000'}, format='json').status_code, 400)
        self.assertEqual(self.client.post(url, {'pin': '123456'}, format='json').status_code, 200)
        self.assertEqual(self.client.delete(url, {'password': 'Str0ng-Pass!'}, format='json').status_code, 200)


class LetterheadRecordTests(APITestCase):
    url = '/api/account/letterhead-records/'

    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        self.user = User.objects.create_user(username='r', email='r@example.com', password='pw-12345!')
        self.client.force_authenticate(self.user)

    def upload(self, content, name='Contract.docx', description='Signed contract'):
        return self.client.post(self.url, {'file': SimpleUploadedFile(name, content), 'description': description},
                                format='multipart')

    def download(self, record_id):
        res = self.client.get(f'{self.url}{record_id}/download/')
        self.assertEqual(res.status_code, 200)
        return b''.join(res.streaming_content), res

    def test_compressible_file_round_trips_identically(self):
        content = b'Letterhead body text. ' * 5000
        res = self.upload(content, name='契約書.docx')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertTrue(res.data['is_compressed'])
        self.assertLess(res.data['stored_size'], res.data['size'])
        self.assertEqual(res.data['sha256'], hashlib.sha256(content).hexdigest())
        self.assertNotIn('file', res.data)
        body, resp = self.download(res.data['id'])
        self.assertEqual(body, content)
        self.assertIn("filename*=utf-8''", resp['Content-Disposition'])

    def test_incompressible_file_is_stored_as_is(self):
        content = os.urandom(4096)
        res = self.upload(content, name='scan.pdf')
        self.assertFalse(res.data['is_compressed'])
        self.assertEqual(res.data['stored_size'], len(content))
        self.assertEqual(self.download(res.data['id'])[0], content)

    def test_validation(self):
        self.assertEqual(self.upload(b'x', description='  ').status_code, 400)
        self.assertEqual(self.upload(b'x', name='virus.exe').status_code, 400)
        self.assertEqual(LetterheadRecord.objects.count(), 0)

    def test_delete_removes_file_and_records_are_private(self):
        rec = LetterheadRecord.objects.get(pk=self.upload(b'hello world' * 100).data['id'])
        path = rec.file.path
        self.assertTrue(os.path.exists(path))
        other = User.objects.create_user(username='o', email='o@example.com', password='pw-12345!')
        self.client.force_authenticate(other)
        self.assertEqual(self.client.get(self.url).json()['results'], [])
        self.assertEqual(self.client.get(f'{self.url}{rec.id}/download/').status_code, 404)
        self.assertEqual(self.client.delete(f'{self.url}{rec.id}/').status_code, 404)
        self.client.force_authenticate(self.user)
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(self.client.delete(f'{self.url}{rec.id}/').status_code, 204)
        self.assertFalse(os.path.exists(path))
        self.assertEqual(self.client.get(self.url).json()['results'], [])

    def test_list_is_paginated(self):
        for i in range(3):
            self.upload(b'doc %d' % i, description=f'Record {i}')
        data = self.client.get(f'{self.url}?page=2&page_size=2').json()
        self.assertEqual(data['count'], 3)
        self.assertEqual([r['description'] for r in data['results']], ['Record 0'])
